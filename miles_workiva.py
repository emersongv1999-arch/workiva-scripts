#!/usr/bin/env python3
"""Pasa a miles los montos de la planilla XBRL bajada de Workiva.

En la planilla conviven hojas en pesos (el balance suma lineas en pesos) y
hojas en miles (las notas linkeadas a lo presentado, x1.000). Para DBNeT un
mismo dato tiene que tener el mismo numero en todos los cuadros, y 1 peso de
diferencia ya es un error: en E211 06-2026 eso fueron 93 de las 115
diferencias que rechazo.

Lo que hace, igual que se arman los estados presentados en miles:
  1. Redondea a miles cada monto escrito (no las formulas), solo en filas
     cuyo elemento XBRL es un monto: tasas, cantidades de acciones, de
     empleados o de clientes, dias y porcentajes quedan como estan.
  2. Recalcula las formulas numericas de la planilla con esos montos, para
     que cada total siga siendo la suma de sus lineas. Redondear el total
     directo no sirve: 4.296.158.573 daria 4.296.159.000, y lo presentado
     (la suma de las lineas redondeadas) es 4.296.160.000.
  3. Escribe los valores nuevos en el mismo .xlsx. Las formulas no se tocan:
     el llenado lee el valor guardado de cada celda.

Solo usa openpyxl para leer; el archivo se reescribe tocando el XML, porque
openpyxl al guardar pierde los valores calculados de las formulas.

APAGADO (REDONDEAR = False): se decidio replicar en Workiva las notas en pesos,
asi todos los cuadros salen de la misma cifra sin redondeo. Mientras este
apagado, ajusta() no toca la planilla y solo revisa el control de la hoja de
obligaciones con bancos.
"""

import io
import re
import sys
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

# Elementos que no son plata. Va por el nombre del elemento porque la
# planilla no trae el tipo de dato.
NO_MONTO = re.compile(r"(NumberOf|PerShare$|Percentage|Proportion|Ratio$|RatioOf|"
                      r"Porcentaje|Numero|Cantidad|Tasa|EnDias$|PeriodoPromedio|"
                      r"Years$|Duration)")
# Pasar a miles. Apagado: las notas se replican en pesos en Workiva.
REDONDEAR = False

# Hoja de obligaciones con bancos: trae su propio control de diferencias.
HOJA_BANCOS = "IFRS7-ObligConBanc"
CONTROL = re.compile(r"^=\+?\$?([A-Z]{1,3})\$?(\d+)-\$?([A-Z]{1,3})\$?(\d+)$")


def es_monto(elemento, valor):
    if NO_MONTO.search(elemento):
        return False
    # ApplicableTaxRate y parecidos: tasas chicas. Un monto que termina en
    # Rate (TaxExpenseIncomeAtApplicableTaxRate) es grande y si se redondea.
    if elemento.endswith("Rate") and abs(valor) < 1000:
        return False
    return True


def a_miles(v):
    """Al mil mas cercano; los ,5 hacia afuera, como ROUND de Excel."""
    return int((Decimal(str(v)) / 1000).quantize(Decimal(1), rounding=ROUND_HALF_UP)) * 1000


def elemento_de_fila(hoja, fila):
    for col in range(1, 13):
        v = hoja.cell(fila, col).value
        if isinstance(v, str) and "#" in v:
            m = re.search(r"#([\w-]+_[\w-]+)", v)
            if m:
                return m.group(1)
    return None


class Planilla:
    def __init__(self, ruta, redondear=True):
        self.formulas_wb = load_workbook(ruta)
        self.valores_wb = load_workbook(ruta, data_only=True)
        self.orig, self.nuevo, self.formulas = {}, {}, {}
        self.redondeadas = 0
        for hoja in self.formulas_wb.worksheets:
            hv = self.valores_wb[hoja.title]
            for fila in hoja.iter_rows():
                elemento, visto = None, False
                for c in fila:
                    v = c.value
                    k = (hoja.title, c.column, c.row)
                    if isinstance(v, str) and v.startswith("="):
                        self.formulas[k] = v
                        self.orig[k] = hv.cell(c.row, c.column).value
                    elif isinstance(v, (int, float)) and not isinstance(v, bool):
                        self.orig[k] = v
                        if not visto:
                            elemento, visto = elemento_de_fila(hoja, c.row), True
                        if redondear and elemento and abs(v) >= 1 and es_monto(elemento, v):
                            nv = a_miles(v)
                            if nv != v:
                                self.nuevo[k] = nv
                                self.redondeadas += 1
        self.calculado = {}

    # ------------------------------------------------------------ formulas
    def valor(self, hoja, col, fila):
        k = (hoja, col, fila)
        if k in self.formulas:
            return self.evalua(k)
        v = self.nuevo.get(k, self.orig.get(k))
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0

    def _rango(self, hoja, c1, f1, c2, f2):
        return sum(self.valor(hoja, c, f) for c in range(c1, c2 + 1)
                   for f in range(f1, f2 + 1))

    TOKEN = re.compile(r"SUM\(([^()]*)\)|(?:'([^']+)'!|([A-Za-z][\w\.\-]*)!)?\$?([A-Z]{1,3})"
                       r"\$?(\d+)|(\d+(?:\.\d+)?)|([\+\-\*/\(\)])|(\s+)")
    REF = re.compile(r"(?:'([^']+)'!|([A-Za-z][\w\.\-]*)!)?\$?([A-Z]{1,3})\$?(\d+)"
                     r"(?::\$?([A-Z]{1,3})\$?(\d+))?")

    def evalua(self, k):
        """Valor de una formula numerica con los montos ya redondeados.

        Entiende lo que usa la planilla para sumar (SUM, + - * /, referencias
        a otras hojas). Cualquier otra cosa (textos, IF, COLUMN...) conserva
        el valor que calculo Workiva."""
        if k in self.calculado:
            return self.calculado[k]
        f = self.formulas[k][1:]
        self.calculado[k] = self.orig.get(k)         # corta referencias circulares
        if re.search(r"[A-Z]{2,}\(", f.replace("SUM(", "")) or '"' in f:
            return self.calculado[k]
        expr, pos = "", 0
        try:
            while pos < len(f):
                m = self.TOKEN.match(f, pos)
                if not m:
                    raise ValueError(f[pos:])
                pos = m.end()
                if m.group(1) is not None:
                    total = 0
                    for parte in m.group(1).split(","):
                        r = self.REF.fullmatch(parte.strip())
                        if not r:
                            raise ValueError(parte)
                        hoja = r.group(1) or r.group(2) or k[0]
                        c1, f1 = column_index_from_string(r.group(3)), int(r.group(4))
                        if r.group(5):
                            total += self._rango(hoja, c1, f1,
                                                 column_index_from_string(r.group(5)),
                                                 int(r.group(6)))
                        else:
                            total += self.valor(hoja, c1, f1)
                    expr += f"({total})"
                elif m.group(4):
                    hoja = m.group(2) or m.group(3) or k[0]
                    expr += f"({self.valor(hoja, column_index_from_string(m.group(4)), int(m.group(5)))})"
                elif m.group(6):
                    expr += m.group(6)
                elif m.group(7):
                    expr += m.group(7)
            res = eval(expr, {"__builtins__": {}}) if expr.strip() else 0  # noqa: S307
            if isinstance(res, float) and res.is_integer():
                res = int(res)
            self.calculado[k] = res
        except Exception:                                # noqa: BLE001
            self.calculado[k] = self.orig.get(k)
        return self.calculado[k]

    def cambios(self):
        """{hoja: {ref: valor nuevo}} de todo lo que cambia."""
        for k in self.formulas:
            self.evalua(k)
        out = {}
        for k, v in list(self.nuevo.items()) + list(self.calculado.items()):
            antes = self.orig.get(k)
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            if not isinstance(antes, (int, float)) or isinstance(antes, bool):
                continue
            if abs(v - antes) > 1e-6:
                out.setdefault(k[0], {})[f"{get_column_letter(k[1])}{k[2]}"] = v
        return out

    def control_bancos(self):
        """Celdas de control de la hoja de obligaciones con bancos que no dan 0.

        La hoja trae al final una resta de dos totales de la misma fila (el
        total de prestamos contra lo que suma la nota). Si no da 0, falta
        algun prestamo o su valor contable."""
        malos = []
        for k, f in self.formulas.items():
            if not k[0].startswith(HOJA_BANCOS):
                continue
            m = CONTROL.match(f.replace(" ", ""))
            if not m or m.group(2) != m.group(4) or m.group(2) != str(k[2]):
                continue
            v = self.evalua(k)
            if isinstance(v, (int, float)) and abs(v) >= 1:
                malos.append((k[0].strip(), f"{get_column_letter(k[1])}{k[2]}", round(v)))
        return sorted(malos)


# ------------------------------------------------------------ escritura

def _rutas_hojas(z):
    rels = {}
    for m in re.finditer(r"<Relationship [^>]*>", z.read("xl/_rels/workbook.xml.rels").decode("utf-8")):
        i = re.search(r'Id="([^"]+)"', m.group(0))
        t = re.search(r'Target="([^"]+)"', m.group(0))
        if i and t:
            rels[i.group(1)] = t.group(1)
    rutas = {}
    for tag in re.findall(r"<sheet [^>]*?>", z.read("xl/workbook.xml").decode("utf-8")):
        nombre = re.search(r'name="([^"]*)"', tag).group(1)
        nombre = (nombre.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                  .replace("&quot;", '"').replace("&apos;", "'"))
        rid = re.search(r'r:id="([^"]+)"', tag).group(1)
        destino = rels[rid].lstrip("/")
        rutas[nombre] = destino if destino.startswith("xl/") else "xl/" + destino
    return rutas


def _numero(v):
    if isinstance(v, int) or float(v).is_integer():
        return str(int(v))
    return format(v, ".15g")


def _parcha(xml, nuevos):
    def cambia(m):
        ref = m.group(1)
        if ref not in nuevos:
            return m.group(0)
        attrs, cuerpo = m.group(2), m.group(3)
        if cuerpo is None or re.search(r'\bt="(s|str|inlineStr|e|b)"', attrs):
            return m.group(0)
        valor = _numero(nuevos[ref])
        if "<v>" in cuerpo:
            cuerpo = re.sub(r"<v>.*?</v>", f"<v>{valor}</v>", cuerpo, count=1, flags=re.S)
        else:
            cuerpo += f"<v>{valor}</v>"
        return f'<c r="{ref}"{attrs}>{cuerpo}</c>'
    return re.sub(r'<c r="([A-Z]{1,3}\d+)"([^>]*?)(?:/>|>(.*?)</c>)', cambia, xml, flags=re.S)


def ajusta(ruta, redondear=None):
    """Deja la planilla en miles (si REDONDEAR) y revisa la hoja de bancos.
    Devuelve lo que hizo, para mostrarlo."""
    ruta = Path(ruta)
    redondear = REDONDEAR if redondear is None else redondear
    p = Planilla(ruta, redondear)
    control = p.control_bancos()
    if not redondear:
        return {"redondeadas": None, "celdas_cambiadas": 0, "control_bancos": control}
    cambios = p.cambios()
    with zipfile.ZipFile(ruta) as z:
        rutas = _rutas_hojas(z)
        destino = {rutas[h]: n for h, n in cambios.items() if h in rutas}
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                datos = z.read(item.filename)
                if item.filename in destino:
                    datos = _parcha(datos.decode("utf-8"), destino[item.filename]).encode("utf-8")
                out.writestr(item, datos)
    temporal = ruta.with_name(ruta.name + ".miles")
    temporal.write_bytes(buf.getvalue())
    temporal.replace(ruta)
    return {"redondeadas": p.redondeadas,
            "celdas_cambiadas": sum(len(v) for v in cambios.values()),
            "control_bancos": control}


if __name__ == "__main__":
    r = ajusta(sys.argv[1], redondear="--miles" in sys.argv or None)
    if r["redondeadas"] is not None:
        print(f"Montos redondeados a miles: {r['redondeadas']}  "
              f"(celdas que cambian, con totales: {r['celdas_cambiadas']})")
    for hoja, celda, v in r["control_bancos"]:
        print(f"  Control de {hoja} no da 0: {celda} = {v:,}".replace(",", "."))
