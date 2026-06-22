"""
verificar_eeff.py — Verificador interactivo de sumas en EE.FF. Workiva

Te pregunta periodo, año, idioma y luego te muestra los documentos
disponibles para que elijas cuáles verificar.

Uso:
    python verificar_eeff.py
"""
import json
import os
import re
import sys
import time
import warnings
from pathlib import Path

from urllib3.exceptions import InsecureRequestWarning
warnings.filterwarnings("ignore", category=InsecureRequestWarning)

try:
    from docx import Document as DocxDocument
except ImportError:
    print("Falta librería: pip install python-docx")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

from workiva_client import get_session, WDESK_BASE

# ---------------------------------------------------------------------------
# Helpers de consola
# ---------------------------------------------------------------------------

MESES = {
    "1": "01", "01": "01", "enero":    "01",
    "2": "02", "02": "02", "febrero":  "02",
    "3": "03", "03": "03", "marzo":    "03",
    "4": "04", "04": "04", "abril":    "04",
    "5": "05", "05": "05", "mayo":     "05",
    "6": "06", "06": "06", "junio":    "06",
    "7": "07", "07": "07", "julio":    "07",
    "8": "08", "08": "08", "agosto":   "08",
    "9": "09", "09": "09", "septiembre": "09",
    "10": "10",             "octubre":  "10",
    "11": "11",             "noviembre": "11",
    "12": "12",             "diciembre": "12",
}

NOMBRE_MES = {
    "01": "Enero", "02": "Febrero", "03": "Marzo",    "04": "Abril",
    "05": "Mayo",  "06": "Junio",   "07": "Julio",    "08": "Agosto",
    "09": "Septiembre", "10": "Octubre", "11": "Noviembre", "12": "Diciembre",
}


def hr(char="─", n=60):
    print(char * n)


def preguntar(prompt, opciones=None, default=None):
    """Lee una línea del usuario. Si opciones, valida que esté en ellas."""
    while True:
        sufijo = ""
        if opciones:
            sufijo = f" [{'/'.join(opciones)}]"
        if default:
            sufijo += f" (Enter = {default})"
        resp = input(f"{prompt}{sufijo}: ").strip()
        if not resp and default:
            return default
        if opciones:
            if resp.upper() in [o.upper() for o in opciones]:
                return resp.upper()
            print(f"  Opción inválida. Elige entre: {', '.join(opciones)}")
        else:
            if resp:
                return resp
            print("  No puede estar vacío.")


def seleccionar_documentos(docs):
    """
    Muestra la lista de documentos y pregunta S/N por cada uno.
    También acepta 'TODOS' o 'NINGUNO'.
    Retorna la lista de documentos seleccionados.
    """
    hr()
    print(f"Se encontraron {len(docs)} documento(s):\n")
    for i, d in enumerate(docs, 1):
        print(f"  {i:>2}. {d['nombre']}")
    hr()
    print("\nOpciones:")
    print("  • Escribe TODOS  → seleccionar todos")
    print("  • Escribe NINGUNO → no seleccionar ninguno")
    print("  • O responde S/N para cada documento\n")

    resp = input("¿Seleccionar todos? [TODOS/NINGUNO/S-N por archivo]: ").strip().upper()

    if resp == "TODOS":
        print("  ✓ Todos seleccionados.")
        return docs

    if resp == "NINGUNO":
        print("  ✗ Ninguno seleccionado.")
        return []

    # Preguntar uno por uno
    seleccionados = []
    print()
    for i, d in enumerate(docs, 1):
        r = preguntar(f"  {i:>2}. {d['nombre']}", opciones=["S", "N"])
        if r == "S":
            seleccionados.append(d)

    return seleccionados


# ---------------------------------------------------------------------------
# Búsqueda en Workiva
# ---------------------------------------------------------------------------

def buscar_documentos(session, mes, anio, idioma):
    """
    Descarga todos los documentos y filtra por mes-año e idioma.
    idioma: 'ESP', 'ENG' o 'AMBOS'
    """
    print("\nConectando a Workiva y descargando catálogo de documentos...")
    all_docs = []
    url = f"{WDESK_BASE}/platform/v1/documents?$top=100"
    while url:
        r = session.get(url, timeout=30)
        r.raise_for_status()
        data = r.json()
        items = data.get("value", data.get("data", []))
        all_docs.extend(items)
        url = data.get("@nextLink") or data.get("nextLink")
    print(f"  Total documentos en workspace: {len(all_docs)}")

    # Patrones de búsqueda
    patron_fecha = [f"{mes}-{anio}", f"{mes}/{anio}"]

    def coincide_idioma(nombre):
        n = nombre.upper()
        if idioma == "AMBOS":
            return True
        if idioma == "ESP":
            return "(ESP)" in n
        if idioma == "ENG":
            return "(ENG)" in n
        return True

    resultado = []
    for d in all_docs:
        nombre = d.get("name", "")
        if any(p in nombre for p in patron_fecha) and coincide_idioma(nombre):
            resultado.append({"id": d["id"], "nombre": nombre})

    resultado.sort(key=lambda x: x["nombre"])
    return resultado


# ---------------------------------------------------------------------------
# Descarga de .docx
# ---------------------------------------------------------------------------

def exportar_y_descargar(session, doc, carpeta: Path):
    """Exporta un documento como .docx y lo descarga. Retorna ruta local."""
    nombre_archivo = re.sub(r'[\\/:*?"<>|]', "-", doc["nombre"]) + ".docx"
    ruta = carpeta / nombre_archivo

    if ruta.exists():
        print(f"    Ya descargado, usando caché: {nombre_archivo}")
        return ruta

    # Iniciar export
    r = session.post(
        f"{WDESK_BASE}/platform/v1/documents/{doc['id']}/export",
        json={"format": "docx"},
        timeout=30,
    )
    if r.status_code != 202:
        raise RuntimeError(f"Export falló {r.status_code}: {r.text[:150]}")

    location = r.headers.get("Location", "")
    # Polling
    for _ in range(40):
        time.sleep(3)
        rp = session.get(location, timeout=30)
        rp.raise_for_status()
        data = rp.json()
        status = data.get("status", "")
        if status == "completed":
            download_url = data.get("resourceUrl", "")
            rd = session.get(download_url, timeout=120)
            rd.raise_for_status()
            ruta.write_bytes(rd.content)
            kb = len(rd.content) // 1024
            print(f"    ✓ Descargado ({kb} KB)")
            return ruta
        if "fail" in status.lower() or "error" in status.lower():
            raise RuntimeError(f"Export error: {data}")

    raise RuntimeError("Timeout en export")


# ---------------------------------------------------------------------------
# Verificación de sumas
# ---------------------------------------------------------------------------

TOTAL_KEYWORDS = [
    "total", "subtotal", "suma", "activos", "pasivos", "patrimonio",
    "resultado", "ganancia", "utilidad", "pérdida", "perdida",
    "flujo", "efectivo",
]


def limpiar_numero(texto: str):
    t = texto.strip().replace("\xa0", "").replace(" ", "")
    if not t or t in ("-", "—", "–"):
        return None
    negativo = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    t = t.replace(".", "").replace(",", ".")
    try:
        valor = float(t)
        return -valor if negativo else valor
    except ValueError:
        return None


def es_fila_total(celdas):
    primera = celdas[0].lower() if celdas else ""
    return any(kw in primera for kw in TOTAL_KEYWORDS)


def analizar_tabla(tabla, tabla_idx, nombre_doc):
    resultados = []
    datos = []
    for fila in tabla.rows:
        datos.append([c.text.strip() for c in fila.cells])
    if not datos:
        return resultados

    num_cols = max(len(f) for f in datos)
    acumulado = [0.0] * num_cols
    tiene_datos = [False] * num_cols

    for i, fila in enumerate(datos):
        while len(fila) < num_cols:
            fila.append("")

        if es_fila_total(fila):
            for col in range(1, num_cols):
                declarado = limpiar_numero(fila[col])
                if declarado is None or not tiene_datos[col]:
                    continue
                diff = abs(declarado - acumulado[col])
                ok = diff < 1
                resultados.append({
                    "documento":   nombre_doc,
                    "tabla":       tabla_idx + 1,
                    "fila":        i + 1,
                    "descripcion": fila[0][:60],
                    "columna":     col + 1,
                    "declarado":   declarado,
                    "calculado":   round(acumulado[col], 2),
                    "diferencia":  round(declarado - acumulado[col], 2),
                    "ok":          ok,
                })
            acumulado = [0.0] * num_cols
            tiene_datos = [False] * num_cols
        else:
            for col in range(1, num_cols):
                v = limpiar_numero(fila[col]) if col < len(fila) else None
                if v is not None:
                    acumulado[col] += v
                    tiene_datos[col] = True

    return resultados


def verificar_docx(ruta: Path):
    resultados = []
    try:
        doc = DocxDocument(ruta)
        for idx, tabla in enumerate(doc.tables):
            resultados.extend(analizar_tabla(tabla, idx, ruta.name))
    except Exception as e:
        resultados.append({
            "documento": ruta.name, "tabla": 0, "fila": 0,
            "descripcion": f"ERROR: {e}", "columna": 0,
            "declarado": None, "calculado": None, "diferencia": None, "ok": False,
        })
    return resultados


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------

def generar_txt(todos, ruta_txt):
    errores   = [r for r in todos if not r["ok"]]
    correctos = [r for r in todos if r["ok"]]
    with open(ruta_txt, "w", encoding="utf-8") as f:
        f.write("=" * 80 + "\n")
        f.write("REPORTE VERIFICACIÓN DE SUMAS — EE.FF. WORKIVA\n")
        f.write("=" * 80 + "\n\n")
        f.write(f"Total verificaciones : {len(todos)}\n")
        f.write(f"Correctas            : {len(correctos)}\n")
        f.write(f"Con diferencia       : {len(errores)}\n\n")
        if errores:
            f.write("=" * 80 + "\nDIFERENCIAS\n" + "=" * 80 + "\n")
            for r in errores:
                if r["diferencia"] is None:
                    f.write(f"\n[ERROR] {r['documento']}\n  {r['descripcion']}\n")
                    continue
                f.write(f"\n{r['documento']}\n")
                f.write(f"  Tabla {r['tabla']}, Fila {r['fila']}, Col {r['columna']}: {r['descripcion']}\n")
                f.write(f"  Declarado : {r['declarado']:>15,.2f}\n")
                f.write(f"  Calculado : {r['calculado']:>15,.2f}\n")
                f.write(f"  Diferencia: {r['diferencia']:>15,.2f}  ◄\n")
        else:
            f.write("Sin diferencias.\n")
    print(f"  Reporte TXT  : {ruta_txt}")


def generar_xlsx(todos, ruta_xlsx):
    if not EXCEL_OK:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verificacion"

    h_fill = PatternFill("solid", fgColor="1F4E79")
    h_font = Font(bold=True, color="FFFFFF")
    ok_fill  = PatternFill("solid", fgColor="C6EFCE")
    err_fill = PatternFill("solid", fgColor="FFC7CE")

    cabeceras = ["Documento", "Tabla", "Fila", "Descripción",
                 "Col", "Declarado", "Calculado", "Diferencia", "Estado"]
    ws.append(cabeceras)
    for cell in ws[1]:
        cell.fill = h_fill
        cell.font = h_font

    for r in todos:
        ws.append([
            r["documento"], r["tabla"], r["fila"], r["descripcion"],
            r["columna"], r["declarado"], r["calculado"], r["diferencia"],
            "OK" if r["ok"] else "REVISAR",
        ])
        fill = ok_fill if r["ok"] else err_fill
        for cell in ws[ws.max_row]:
            cell.fill = fill

    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["D"].width = 40
    for c in ["B","C","E"]: ws.column_dimensions[c].width = 8
    for c in ["F","G","H"]: ws.column_dimensions[c].width = 16
    ws.column_dimensions["I"].width = 10

    ws2 = wb.create_sheet("Resumen")
    ws2.append(["Documento", "Total", "OK", "Diferencias"])
    for cell in ws2[1]:
        cell.fill = h_fill
        cell.font = h_font
    docs = {}
    for r in todos:
        d = r["documento"]
        docs.setdefault(d, {"total": 0, "ok": 0, "err": 0})
        docs[d]["total"] += 1
        if r["ok"]: docs[d]["ok"] += 1
        else:        docs[d]["err"] += 1
    for nombre, cnt in sorted(docs.items()):
        ws2.append([nombre, cnt["total"], cnt["ok"], cnt["err"]])
        for cell in ws2[ws2.max_row]:
            cell.fill = ok_fill if cnt["err"] == 0 else err_fill
    ws2.column_dimensions["A"].width = 55

    wb.save(ruta_xlsx)
    print(f"  Reporte Excel: {ruta_xlsx}")


# ---------------------------------------------------------------------------
# Flujo principal
# ---------------------------------------------------------------------------

def main():
    hr("═")
    print("  VERIFICADOR DE SUMAS — EE.FF. WORKIVA")
    hr("═")

    # --- 1. Periodo ---
    print("\n📅  PERIODO\n")
    mes_raw = preguntar("  Mes (número o nombre, ej: 6 / junio / 06)")
    mes = MESES.get(mes_raw.lower())
    while not mes:
        print(f"  Mes '{mes_raw}' no reconocido.")
        mes_raw = preguntar("  Mes")
        mes = MESES.get(mes_raw.lower())

    anio = preguntar("  Año (ej: 2026)")
    while not re.fullmatch(r"\d{4}", anio):
        print("  Año inválido.")
        anio = preguntar("  Año")

    # --- 2. Idioma ---
    print("\n🌐  IDIOMA\n")
    idioma = preguntar("  Idioma", opciones=["ESP", "ENG", "AMBOS"], default="AMBOS")

    # --- 3. Buscar en Workiva ---
    print(f"\n🔍  Buscando documentos {mes}-{anio} [{idioma}]...")
    session = get_session()
    docs = buscar_documentos(session, mes, anio, idioma)

    if not docs:
        print(f"\n  No se encontraron documentos para {mes}-{anio} [{idioma}].")
        sys.exit(0)

    # --- 4. Selección interactiva ---
    print()
    seleccionados = seleccionar_documentos(docs)

    if not seleccionados:
        print("\n  Ningún documento seleccionado. Saliendo.")
        sys.exit(0)

    # --- 5. Descargar ---
    nombre_carpeta = f"docx_{mes}_{anio}_{idioma}"
    carpeta = Path(nombre_carpeta)
    carpeta.mkdir(exist_ok=True)

    hr()
    print(f"⬇️   Descargando {len(seleccionados)} documento(s) en ./{nombre_carpeta}/\n")
    rutas = []
    for i, doc in enumerate(seleccionados, 1):
        print(f"  [{i}/{len(seleccionados)}] {doc['nombre']}")
        try:
            ruta = exportar_y_descargar(session, doc, carpeta)
            rutas.append(ruta)
        except Exception as e:
            print(f"    ✗ Error: {e}")

    # --- 6. Verificar sumas ---
    hr()
    print(f"🔢  Verificando sumas en {len(rutas)} archivo(s)...\n")
    todos = []
    for i, ruta in enumerate(rutas, 1):
        print(f"  [{i}/{len(rutas)}] {ruta.name}")
        res = verificar_docx(ruta)
        ok  = sum(1 for r in res if r["ok"])
        err = len(res) - ok
        print(f"    Verificaciones: {len(res)} | OK: {ok} | Diferencias: {err}")
        todos.extend(res)

    # --- 7. Reportes ---
    hr()
    total_err = sum(1 for r in todos if not r["ok"])
    print(f"\n📊  RESULTADO FINAL")
    print(f"  Total verificaciones : {len(todos)}")
    print(f"  Correctas            : {len(todos) - total_err}")
    print(f"  Con diferencia       : {total_err}")

    prefijo = f"reporte_{mes}_{anio}_{idioma}"
    generar_txt(todos, prefijo + ".txt")
    generar_xlsx(todos, prefijo + ".xlsx")

    hr("═")
    print("  Listo.")
    hr("═")


if __name__ == "__main__":
    main()
