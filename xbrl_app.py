#!/usr/bin/env python3
"""Enlace XBRL: llena las plantillas de DBNeT directo desde Workiva.

Se elige empresa y periodo; la aplicacion descarga la planilla
"E___ XBRL MM-AAAA" desde Workiva, revisa, llena los .xlsm de DBNeT de esa
empresa en su lugar y arma el archivo unico con los botones funcionando.

Cada empresa es una carpeta dentro de la carpeta de trabajo (la que se
elige arriba a la derecha y la app recuerda), con sus plantillas en xls\\:

    <carpeta de trabajo>\\
        E211\\xls\\   las plantillas .xlsm de DBNeT de E211
        E205\\xls\\   ...

Las plantillas se cargan desde la app. Una empresa aparece sola en cuanto
existe su planilla XBRL en Workiva.

El llenado y la fusion corren como procesos aparte -- el mismo .exe llamado
con --tarea -- y su salida se lee linea a linea: la ventana sigue
respondiendo, y si Excel se cae armando el archivo unico no se lleva la
ventana con el.
"""

import io
import json
import os
import queue
import re
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import zipfile
from pathlib import Path

RESPALDO = "_original_dbnet"


def carpeta_base():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


# ════════════════════════════════════════════════════════ tareas (hijo) ══
def corre_tarea(nombre, args):
    """El .exe llamado con --tarea: corre el llenado o la fusion y termina."""
    salida_alt = None
    if sys.stdout is None:
        # Un .exe sin consola a veces arranca sin salida estandar aunque el
        # padre le pase una tuberia; entonces se escribe a un archivo que el
        # padre lee al terminar.
        salida_alt = open(os.environ.get("XBRL_SALIDA", os.devnull), "w",
                          encoding="utf-8", buffering=1)
        sys.stdout = sys.stderr = salida_alt
    else:
        for flujo in (sys.stdout, sys.stderr):
            try:
                flujo.reconfigure(encoding="utf-8", errors="replace",
                                  line_buffering=True)
            except (AttributeError, ValueError):
                pass

    if nombre == "llenar":
        import llenar_dbnet_desde_workiva as modulo
    elif nombre == "fusionar":
        import fusionar_cuadros as modulo
    else:
        print(f"Tarea desconocida: {nombre}")
        return 2

    sys.argv = [nombre, *args]
    try:
        modulo.main()
        codigo = 0
    except SystemExit as e:
        if e.code is None or isinstance(e.code, int):
            codigo = e.code or 0
        else:
            print(e.code)
            codigo = 1
    except Exception:                                    # noqa: BLE001
        traceback.print_exc()
        codigo = 1
    sys.stdout.flush()
    if salida_alt:
        salida_alt.close()
    return codigo


# ═══════════════════════════════════════════════════════════ ventana ══
if __name__ == "__main__" and len(sys.argv) > 2 and sys.argv[1] == "--tarea":
    sys.exit(corre_tarea(sys.argv[2], sys.argv[3:]))


import tkinter as tk                                     # noqa: E402
from tkinter import filedialog, messagebox, ttk          # noqa: E402

from xbrl_workiva import ErrorWorkiva, Planilla, Workiva  # noqa: E402

TEMPORAL = Path(tempfile.gettempdir()) / "XBRL_llenado"
# La carpeta interna conserva el nombre de antes del cambio a "Enlace XBRL":
# ahi estan la carpeta de trabajo elegida y la lista de planillas guardada.
DATOS = Path(os.environ.get("LOCALAPPDATA") or tempfile.gettempdir()) / "XBRL_DBNeT"
CACHE = DATOS / "planillas.json"
CONFIG = DATOS / "config.json"
ULTIMO = "ultimo_llenado.json"


def nota_ultimo(carpeta):
    """Donde la app anota el ultimo llenado de una empresa.

    Dentro de xls\\_original_dbnet, la carpeta interna del respaldo, y no a
    la vista en la carpeta de la empresa, donde un .json confundia. Ademas
    asi viaja con las plantillas: al cambiarlas por otras, la carpeta xls
    entera pasa a plantillas_anteriores y la nota se va con ella."""
    return carpeta / "xls" / RESPALDO / ULTIMO
# El icono (un eslabon) va embebido para no depender de un archivo al lado
# del .exe; el .ico del ejecutable lo pone COMPILAR_XBRL.bat (enlace_xbrl.ico).
ICONO_64 = "iVBORw0KGgoAAAANSUhEUgAAAEAAAABACAYAAACqaXHeAAAM30lEQVR4nNWbe3Bc1X3HP79z7933SvL7hXnUxAEbMIamJSGxTaHFwQHKQ04hGQgzCdQYSkOblCGTykqmTdMBQsHQQpi0k2CaWDSJS6AQAnGgpJ22A3aIY3CpB+MH4Jfs1e5qd+8959c/7kqWjSSvQGvJ35mVVtp7zt3P75zfOb9z7vkJI1VHh2E9BpZAcYuMuHwzlJursB6W4OjsVEBH+Q4qLO7wR7nS5mlxhw/aUOMc/aL2tR5dy23fn/lPPzxXNTpTVU/C0QLuA3zT0ZBBoWjE34YX/br4g5Wb+z864rsPpuEMIHR0CJ2drnXZgxOijPc51F0LboH4qQAxo0UwOlKHhhWLMa8q+gOJ+E7pRzfvpqPD0LlKQQZ1iyEMoNJXIHP1AzeIeJ3iJ2erDcHWQNUig1c4ZlIVRDy8BOIFuKjyrlH3teLjKx+ML+gw0Pme7jqIAerw7R2JrJv2bUmkr9OoAjaKQAyCDF5uXEhRFNRhfF+CFBpW/qUk73yOrs7iYEY4AkSF9uWGno/72Zy/ThK5i7VaiFAxCOOszx9FqgpEksoHWiu/lBG3dM+8PeUj3eFwqPYuQ1eXzWTMI5LIXayVQg3EP+7gAUQEkUArhVASmfPL1q2hs9PR3nUYy6E/6iNm5or7rjPJ3Ge1WggRSRzzLz6MBBABI/HvxgqZQCs9oaRaLstecf+tdC23tK/1BtZJ35zZevm9raEfvCbGn4KNGC8tb+rkUaREVlFVjBF8L345VXT4IdkhHqgrqtEPl7tueTeOlURjwMWrPBCNfO8Gk8hOw4Vu3MAboVJz9JQi0inDybNSnPZbGWZNTeIZKBQjoig2yHDVoJGTRKZFHCsAjZkH9gCF7FWrX8ZPLiCqOQRvmAqbLqn/KJYtp52S4VNLJnHW3ByteR/fE2qhY293yH+92sMT6/fy7r6QXMbDuSG6guLwAyGqbS1NTpzOwzdFAH4cKIhrWb76VCvmDGxNxrr1+/y73Ou45pKpfPZT00mlDNWqI7KxG/iecMK0JKeemOai8yZw/2M7eOnlg+Sy/uBGEAy2phh/TmZf5YwyvEL7Wq++sIEw4mwJUj6qljGc5/vgi2XL56+ewU3LZ2GdUuiJCCPFOVCnWAe1UOkuROSzHn+54mQ+fm4bxbIdxh3Eip8E9c8BYPcmiVd1gKAnI4axjPAGwt/YPpNrLplKdyFEFTxPUCCdNOQyPslAUKXfHaJIufXaWcyYnKAWuqFnCRFE7CnxjWaK6VvSitH8MWAcUoPBH+yJ8Aa0ZiZp2PJWmZdePsDbe2tk0wbVQwPllIkBl//eZKpVhxlmnlSVmDU3VwcsccdXyx/siTAmbmUjEASGB76/k397cT/OKqmk4fNXz2DZokmUei2eEXqrjt85M89jT/mEoQ7XC4aIBMdAR4MXgWTCcO/3trPu2T0kAiGT9git8tDaXby5q0IyYVAgipRJbQHTJiYIo2HcYIDGxWg/HHwqYbjnu9t5+oX9tLQGqIJ1SjJhqNaUrTsqJAID9WAoERjyWQ/rGosWx8wAI4Zv8bH2kJeqAgITW31s37QnYK1Srbk4emzAqcfEAB8U3veEnkLIko+0ceaHsvRWHCKCb6CnFLGnO8T3paGNwWNugA8K73nCwZ6I83+7jduvP4Go/plTJZn02Ly1zJ79IYFvjrY+AI6xAUYDvtATcf45rdx540mAYG1cTkSwVvnX9fvi+zRCzzE0wGjDq0Jk46nOWmVCi8/jP93Dhs09ZNIeQy0JjtQxNIBQ7nWDwvuekE173PPdHSOCNxLPCG0tPuue38s//uhtsulhFkSDqOl7/ULdbw/WuP7KmVy7bBrdhbA/whOB3jDkrx7Zzi//p0hba0AYHdq2Oyp83ueJ9fv4u+/tIJU08Q1HENI1tQdIfefm4MEal3xixmEt3+emimNey2mcPWMOiSRUqlF/GDtSeJGGXb9fTTOACAhCT6nK335hEX9983zK1RCpbysr4LCcPXkep0+dzQM3L2bdVy8jk/QJIxvDF5oLD00yQD98scJ9Ky7kS1edw7buPXjigSiCEDrLgonzODE/A6eOMHIsPfdkHr/zUgCqVceij7Rx5xeaBw9NMMBA+G/98QXceukCdvTsJtIIIwYRoWpDPjRxJqe0zayXMQS+oRZZLjjrBG69dCE1G7Li07PwvHgvsBnwMMoGGAh/74oL+NPLF+Kc4jj0eM45JZPyeHHDXh58YmPdLajPBganyhevWEjKT/LqG0V8L3aYZsDDKBrgSPjbLltIaOPdGV+8/pHZqZIKfF5/Zzcrv/4USzt+TKFcA7Q+OsCMCVnOPXUar28rEvhCZGkKPIySAQaDj6zD9+Lq88ksqSDAqeIZoVi2/P5HJ/OJi0/gxf/Yypr1r8WRnHM4F6/qTp89kUIxxDloa2kOPIyCAYaFrz+iz3gZahUPYxSU/sivY8VJzJ6TZ9O2/e+pszWbxPbD720KPHxAAwwHr9r3eA6uv/tZ7nr0DfLZAOsUp0o2ZXj6pW62v1nivNOmx/UNqPvd7jKtubjl72sSPHyASLAReGOEG1c/x5qfvUqQSvLUC/tYtmgSkVWefLGbu1a/zvJPnsEfLZqLU8V4pp9w6+5u3u4p8fP/7GkaPLxPAzQMf//P+PYTG2ibmKUWOu5fs4N00lAL4a77X+fqi+ex5stL8eqHLZyLp7vfvLWPjVvfJbTxOsHzpCnw8D4MMFL41glZwshhDFgLd//TDoqFXtqXzuef/+KSetiriAjOKZ4xfKPrvymWa7Tm00TWNQ0eRjgGSP151UjgIxsvbFQhmfAoFntZfuHpfP+OSzASn2gQEULrCDzDd57dxJrnN5PPpfrLNlMjMwBCsVQdMTxA4BsOdJdZfsHpPPalT/avB8wA+Ed//ho33vcs2XSifwBttho2gOcZeooVvnLN79aDnPcPH7f8EfDrX+P6u58mlQxAmufzR6ohAxgRyr01Fs6dzqrPnId1im/q8IwA/svDwN/1NOlUgBE5Zq0PjRrACFHN8ieXnl0H1/q0pBgRblr93KjBu2MIDw0YQARqkaWtLc2FC2bHACYesY0Rbn/kBR5e9zJtE48/eGjEAAhh5Jg5Mcf0Cdn+lZsxwpad3Tzwk43k2jJE0fEHDw26gFMlnfAI/MMv37W/RGgdnjH923DHEzw0YAAFPCMcKFUpV0MgdgsF5p84iSktKXrKNXzPYIxwYF/puIGHRgygSsL32Lm3yNZ3CvEIrfEpjSmtaR665SIm51MUSlV6KyE3LDurDi/jHh4OC4WHPl7ue4ZSsULXv/8vZ3zmPKzriwGUP/zoHD42byavbt3N5LYMC06ZEtem4xheD7GaONkAVE1xqOutc6QySf7hyY3s3FfE9wzWuv74fWprmgsXnsSCU6agqvEJPIEwGofwgBhi1uIWMbA+/qewDXWHWadPqpDwDXsO9nL9Pc/EA59nCK0DiQdJ61z/Y2pXfx/4cXg7nuBBUcc2AHK71LAkzngwxmzUsGIR8Rjk2Yp1Sj6X5LlX3uKyr61jx94igWcwIvHLGDwjiAhe/f29617hhm89M57gPY2qqNpXAJg6v6+1VehYJZlfT94oXnL+cAclPc9QKFaYOSnHzcsWcOXHTmXO9FYSgYcCew6UeXHTTh588lc8v2EbuUyyHtuPMbzi8Hwhqr1V6tUP8/RtVfqf2S7u8PlFZ5S58r47TDL/Da0VI5Ah9wo8I1RDS7U3JJdPceKUPBNySSo1y679Rd7eVwQRWjKJeJPzWEEOK40kkfO1Wvpm6Ye33NHHHE+Dv1hlATG12iOuVtqPBIZhkoFi//ZobU3jVNmys5tfbt7FK1t3s7+nQks2SUsmgR038CjiGQ3LZWvlAUD6XT/+XJT2tab4kz/fa9CvSCJl6idGh65Rlcg6DJBO+OTSCbLJgMA3WKeHzu2MB6lGksgadfbrlXUrt9O+1tAZZ44cPuL35QxctfqHJpm7QiuFEDHBmHzp0ZK6UJL5wFWLz5XP3PsH/Ga+0LXcUR/o35sy07FK+L856Wyl+IwkMudrpSdE8EeQojBOpIoSSTIXaNS7waP3wkLX7d2wSgbmDR0RCovSCTx6XakkbqnWyj+WZD5APEE1Qsc8SfDoUhyqEWJEUvlAo8pPTXjgokLXn+2n43B4GPJU+KHsqmz7398G8lXxk5PU1uK0OSRqOsj7knqYQMRPorZ2UF30zfLjK/8G0DgtoKG0uf7KpK+7pNvvnmXI3gi6HDhN/OTwRcdEikY1QN8A83hUKz9UXffFNwHpS48ZrNTRKQamny7u8DPTZp6F2rNF5WRV2zJ6AO9HDjAIpihGtlkxG3oLlV/Vg5yGUmcb1HGYPN3R0dBmz0j7sdC+1rB7k/QlWowfrYep83XgFNeI/h9Kn6ogQy5DhwAAAABJRU5ErkJggg=="
ICONO_32 = "iVBORw0KGgoAAAANSUhEUgAAACAAAAAgCAYAAABzenr0AAAF6klEQVR4nMWXf2xVZxnHP8/7nnPvuT/ovf3FRrktMBmaEjALm8sWo7Mm4jK3sCZ3soT/piOR1BmW7I/paJqOaeI/CkuYbMaY+YelQ9hiEIJLxrag4gIixoyBlNqC/UFLb9v7o/ee877+cduuLUXbuxmf5CQn58f7+X6f87w5zwNzo71d8b+OBQyZd6Ojw0Bax9Mtj1nhaxizBqy2Zs5zywgBi5IAUX0CJyeHBo5yqsP/mDUjYPpC9PF9nxcn/Jo47r3l103lTucpKZu2/tRfzFTu2/m3dn8ww5SZk1h6/yZw3xXtJG0p75cF2OU7F1BKZlNmDFiwWIu4Eccaf8L6pYdyR9rOkj6kywL+1aCjN4ofqFBksy3mS4i4lRhVAr6BfD6Yvea6ghfWGGPB2pK4Edf4hQ9zN4bv4RRFAYi07tumw/Ejtpj1QZxK4FoJ+amAWETzpXuTNK0KM5kNOH1+nMu9eaIRhbWAtb6E4471c9uz3bu6HACFegSwVFZraC1kcwGNd4Z5/uk1rG+K4DpCEFjSW1dyoOsax94dJRbVGANgLcY8AnSVq0NoxFqp5JvPwFN3hulsW8e6lEeuYPjbpSyDIyW0Fr6zfTUb1kYoTAUICMaItdJYNl/WU9H+nwt/sW0dK2tCmABef2uA7+69xEuv9jI24RP1FA/ek6BUsqgZklg1K+DTgNfXhEDg4BvX6To2hOdpzl+cZHi0iNaKqpgGwC5YpzLn6la4CBzsvs7hE8NUJx0mJn22fbWOtas9fN/QNzBVNr5grWVXvKOFbD4gdUcZfkddGCXwyqEyPJl0uDnm82hLHTvTDThauHq9wDtnxoh400U4x/ayBDhaMZ4tsqo+zN5n7qKu2iFbKPHKoWuceOcmyaTL2JjPoy217HpyNQDDN0vsPdhLZtInElYYY+dlYcmfwNGKTHaKuxuq+VHb56ivcRjPFdlcu4GWz2wiElNkJsrOdz2ZAmA0U2LP/h6u9BWIehqzsACWmgFHKzKTBTatqefNjofpDT5kMDPJA6lmmmKNPJduJLCG7j+f43s7UpQCy2imxAv7e+jpLxCPaYJgEfpSMlCGT9HcVMvxzlYa6qLcGM9THQ+z91cX+P25fgJjeXrrJj6bqsaKWTL8vwqYgW9squFEZysNtVGUdYm4IbRSrEzl+eEbf0QrIRkP0VC9gr7BAu0vX10S/D8KmAs/3tlKqi4OwLM/O80vf9tPTZXLR1fzfLE5hbUWYy39wzk6D/Tyj778kuBwmxrQt4E/94v32f+bM1TXxbl49QIP3nU3L2y/HxHhfM8gR079E6UVsejS4ItmQCkhly/S3FTD8Rfnwt/jx7/+E7X1K7g5muULTes5sOsrOFoYyuR56qcnQSzhUHmrLTXm92ci+IGhKhrmzT2PkaqdCz9DbW2MkZFJ2lq38OozLQAMZXI8vOcIf+0ZJhIOLdn5ogK0FvLZIt/88gbWr0oSGMv3Xz89DY8zMpKlrXUL+3Y+NAd+lLOXBknEPQKz/BZukSK0bGysxVhLz2CGl7rOEKvyZp0vBk+u8PCDyvrHWwVYuDGeR4mwMhHl61vWUsoW2f3EfbeFlyqEw/QukHKzhAksjudy+PRlfrD9fqqiIQ4//w36h8fZkKr5dOG2/Gee7oikDxFrwEbDDheuDLPz5bcplAKiYWcWfmUg88nhIhZRFuifzYCx9ncaeQosgYEV8TA/P3GBs5cH2fbAeuqqPD66NkbXexcZGsuR+ETOLYiIaH0MQKBdkd7oRIPBc8qNNNtSuS3XSsgWSgRFn5k2wouGcJ3l7fMFcB/Xc2ypcDmni5vp3l1QpDcK3U8UEbPDGn9CXM/FWj8w1o+GnCCRiASJhBckEl7gKAlMYAKsreTwxfUcrMlZkR10P5sn3a3mjWaRx39ynzjea8oJby7PM+bWJq6SUAoQbFD8e1AsfqtwtO0PH49mMzEzMKbbQzFpaMXardYETWA1VDCiASAWJBBRfShOZgcHDnOqo0A6renuDm59/v8wnv8bYf8HkrrG+REAAAAASUVORK5CYII="
SIN_CONSOLA = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0

AZUL, AZUL_OSC, FONDO, BLANCO = "#0B5394", "#083D6E", "#F4F6F8", "#FFFFFF"
VERDE, ROJO, GRIS, TENUE = "#1E7B34", "#B3261E", "#5A5A5A", "#8A94A0"
AMBAR, AMBAR_FONDO = "#8A5A00", "#FFF4DC"
FUENTE = "Segoe UI"

# "  IAS12-Cuadros(835110)_2025.xlsm    202 celdas  (dry-run)", una por
# plantilla. El nombre va en un campo de 50 de ancho: eso la distingue del
# "OJO: N avisos", y aguanta los nombres largos a los que el script les come
# la extension.
AVANCE = re.compile(r"^\s{2}.{50}\s*\d+\s+celdas\b")
N_PLANTILLAS = re.compile(r"^Plantillas:\s*(\d+)")
TOTAL = re.compile(r"^Total:\s*([\d.,]+)\s+celdas\s+en\s+(\d+)\s+hojas")
AVISOS = re.compile(r"OJO:\s*(\d+)\s+avisos")
EMPRESA = re.compile(r"E\d+", re.IGNORECASE)

# Lo que se esta haciendo, dicho para quien usa la app y no para quien la
# programo: «simular» o «armar el archivo unico» no le decian nada a nadie.
PASOS = ("Trayendo los datos desde Workiva…",
         "Revisando que los datos calcen con las plantillas…",
         "Llenando las plantillas…",
         "Preparando el archivo para DBNeT (toma un par de minutos)…")
FALLO = ("traer los datos desde Workiva",
         "revisar los datos",
         "llenar las plantillas",
         "preparar el archivo para DBNeT")


# ═══════════════════════════════════════════════════ carpeta de trabajo ══
def lee_config():
    try:
        return json.loads(CONFIG.read_text("utf-8"))
    except (OSError, ValueError):
        return {}


def guarda_config(datos):
    try:
        DATOS.mkdir(parents=True, exist_ok=True)
        CONFIG.write_text(json.dumps(datos, indent=1), "utf-8")
    except OSError:
        pass


def carpeta_de_trabajo():
    """La carpeta con las empresas. Se recuerda entre una vez y otra, asi
    que da lo mismo desde donde se abra el .exe: antes se buscaba junto al
    .exe, y al moverlo las plantillas cargadas parecian perdidas."""
    guardada = lee_config().get("carpeta")
    if guardada and Path(guardada).is_dir():
        return Path(guardada)
    carpeta = carpeta_base()
    guarda_config({**lee_config(), "carpeta": str(carpeta)})
    return carpeta


BASE = carpeta_de_trabajo()


def plantillas_de(carpeta):
    xls = carpeta / "xls"
    if not xls.is_dir():
        return []
    return [p for p in xls.rglob("*.xlsm")
            if RESPALDO not in p.parts and not p.name.startswith("~$")]


def empresas_locales():
    if not BASE.is_dir():
        return []
    return sorted(d.name.upper() for d in BASE.iterdir()
                  if d.is_dir() and EMPRESA.fullmatch(d.name))


# ═══════════════════════════════════════════════════ plantillas de DBNeT ══
def reune_plantillas(origen):
    """Las plantillas .xlsm que hay en la carpeta origen (y sus subcarpetas),
    sueltas o dentro de un .zip como el que entrega DBNeT.
    Devuelve ({nombre: bytes}, descartes, repetidas, ya_de_la_app).

    Se exige que cada una sea un .xlsm con macros: lo que DBNeT entrega. Un
    .xlsx o una copia rota se descarta con su motivo en vez de terminar en
    xls y fallar a mitad del llenado. Lo que ya esta dentro de una carpeta
    de empresa de la app (llenado, o respaldo) se salta: cargarlo de nuevo
    mezclaria plantillas llenas con virgenes."""
    halladas, descartes, propias = {}, [], 0
    donde = {}                  # nombre -> [rutas], para explicar repetidas
    repetidas = set()
    origen = Path(origen)
    empresas = [BASE / e for e in empresas_locales()]

    def de_la_app(ruta):
        r = ruta.resolve()
        return any(r == e.resolve() or e.resolve() in r.parents for e in empresas)

    def considera(nombre, leer, desde):
        partes = nombre.replace("\\", "/").split("/")
        base = partes[-1]
        if (not base.lower().endswith(".xlsm") or base.startswith("~$")
                or RESPALDO in partes or "plantillas_anteriores" in partes
                or "__MACOSX" in partes):
            return
        try:
            contenido = leer()
            with zipfile.ZipFile(io.BytesIO(contenido)) as z:
                nombres = set(z.namelist())
        except (OSError, zipfile.BadZipFile) as e:
            descartes.append((base, f"no se pudo leer ({e})"))
            return
        if "xl/workbook.xml" not in nombres:
            descartes.append((base, "no es un libro de Excel"))
        elif "xl/vbaProject.bin" not in nombres:
            descartes.append((base, "no trae las macros de DBNeT"))
        else:
            donde.setdefault(base, []).append(desde)
            if base in halladas and halladas[base] != contenido:
                # Dos archivos con el mismo nombre y distinto contenido: uno
                # puede ser una copia ya llenada. Elegir cualquiera seria
                # adivinar, asi que no se carga nada hasta que se aclare.
                repetidas.add(base)
            halladas.setdefault(base, contenido)

    archivos = [origen] if origen.is_file() else sorted(
        p for p in origen.rglob("*")
        if p.is_file() and p.suffix.lower() in (".xlsm", ".zip"))
    for ruta in archivos:
        if de_la_app(ruta):
            propias += 1
            continue
        relativa = str(ruta.relative_to(origen)) if origen.is_dir() else ruta.name
        if ruta.suffix.lower() == ".zip":
            try:
                with zipfile.ZipFile(ruta) as z:
                    for info in z.infolist():
                        if not info.is_dir():
                            considera(info.filename, lambda i=info, z=z: z.read(i),
                                      f"{ruta.name} → {info.filename}")
            except (OSError, zipfile.BadZipFile) as e:
                descartes.append((ruta.name, f"no se pudo abrir el .zip ({e})"))
        else:
            considera(relativa, ruta.read_bytes, relativa)
    repetidas = {n: donde[n] for n in sorted(repetidas)}
    return halladas, descartes, repetidas, propias


def instala_plantillas(carpeta, plantillas):
    """Deja plantillas como el juego completo de la empresa en carpeta\\xls.

    El juego anterior no se borra ni se mezcla: la carpeta xls entera, con
    su respaldo _original_dbnet, pasa a plantillas_anteriores\\<fecha>.
    Mezclar no sirve: cuando DBNeT cambia el ano de una plantilla cambia su
    nombre (..._2025.xlsm pasa a ..._2026.xlsm), y con las dos en xls el
    cuadro se llenaria dos veces. Devuelve donde quedo el juego anterior."""
    xls = carpeta / "xls"
    archivado = None
    if xls.is_dir() and any(xls.iterdir()):
        archivado = carpeta / "plantillas_anteriores" / time.strftime("%Y-%m-%d_%H%M%S")
        archivado.mkdir(parents=True)
        try:
            xls.rename(archivado / "xls")
        except OSError as e:
            archivado.rmdir()
            raise PermissionError(
                f"No se pudo apartar las plantillas actuales de {carpeta.name}: "
                "seguramente hay alguna abierta en Excel. Cierra Excel y vuelve "
                f"a intentar.\n\nDetalle: {e}") from e
    xls.mkdir(parents=True)
    for nombre, contenido in plantillas.items():
        (xls / nombre).write_bytes(contenido)
    # Las macros de cada plantilla escriben sus CSV en la carpeta csv que
    # esta al lado de xls, y no la crean.
    (carpeta / "csv").mkdir(exist_ok=True)
    # Las plantillas nuevas estan virgenes: lo que decia el ultimo llenado
    # ya no es cierto.
    nota_ultimo(carpeta).unlink(missing_ok=True)
    (carpeta / ULTIMO).unlink(missing_ok=True)    # donde la dejaban versiones anteriores
    return archivado


def lee_ultimo(carpeta):
    try:
        vieja = carpeta / ULTIMO
        if vieja.exists():
            # Versiones anteriores la dejaban a la vista: se lleva adentro.
            nota_ultimo(carpeta).parent.mkdir(parents=True, exist_ok=True)
            vieja.replace(nota_ultimo(carpeta))
        return json.loads(nota_ultimo(carpeta).read_text("utf-8"))
    except (OSError, ValueError):
        return None


def miles(n):
    return f"{n:,}".replace(",", ".")


def plural(n, uno, varios):
    return f"{n} {uno if n == 1 else varios}"


def abre(ruta):
    if os.name == "nt":
        os.startfile(ruta)                               # noqa: S606
    else:
        subprocess.Popen(["xdg-open", str(ruta)])


def muestra_en_carpeta(ruta):
    """Abre la carpeta con el archivo ya seleccionado."""
    if os.name == "nt":
        subprocess.Popen(["explorer", "/select,", str(ruta)])
    else:
        abre(Path(ruta).parent)


# ═══════════════════════════════════════════════════════════ ventana ══
class Aplicacion(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Enlace XBRL")
        self._pon_icono()
        self.configure(bg=FONDO)
        # Los notebooks de CGE van con la pantalla al 125-150%: el tamano se
        # escala con eso y se limita a la pantalla, o la ventana no cabe.
        self.escala = escala = max(1.0, self.winfo_fpixels("1i") / 96)
        ancho = min(int(1000 * escala), int(self.winfo_screenwidth() * 0.95))
        alto = min(int(720 * escala), int(self.winfo_screenheight() * 0.88))
        self.geometry(f"{ancho}x{alto}")
        self.minsize(min(ancho, int(820 * escala)), min(alto, int(580 * escala)))

        self.empresa = tk.StringVar()
        self.periodo = tk.StringVar()
        self.info_workiva = tk.StringVar()
        self.info_plantillas = tk.StringVar()
        self.info_carpeta = tk.StringVar()
        self.estado = tk.StringVar()
        self.planillas = self._lee_cache()
        self.cola = queue.Queue()
        self.confirmacion = queue.Queue()
        self.corriendo = False
        self.buscando = False
        self.hechos = self.total = self.paso = 0
        self.registro = []          # todo lo que imprimieron los scripts
        self.ventana_log = None

        self._estilos()
        self._construye()
        self._refresca_carpeta()
        self._refresca_empresas()
        self._busca_planillas()
        self.after(80, self._drena)

    # ------------------------------------------------------------ aspecto
    def _pon_icono(self):
        """El eslabon en la barra de titulo y de tareas, en vez de la pluma
        de tkinter (que es tambien la del Auditor)."""
        try:
            self._iconos = [tk.PhotoImage(data=d) for d in (ICONO_64, ICONO_32)]
            self.iconphoto(True, *self._iconos)
        except tk.TclError:
            pass

    def _estilos(self):
        s = ttk.Style(self)
        try:
            s.theme_use("vista" if os.name == "nt" else "clam")
        except tk.TclError:
            pass
        s.configure("TCombobox", padding=5)
        s.configure("Treeview", rowheight=int(24 * self.escala), font=(FUENTE, 9))
        s.configure("Treeview.Heading", font=(FUENTE, 9, "bold"))
        s.configure("Horizontal.TProgressbar", thickness=8)

    def _construye(self):
        cab = tk.Frame(self, bg=AZUL, padx=24, pady=10)
        cab.pack(fill="x")
        izq = tk.Frame(cab, bg=AZUL)
        izq.pack(side="left", fill="x", expand=True)
        tk.Label(izq, text="Enlace XBRL", bg=AZUL, fg=BLANCO,
                 font=(FUENTE, 18, "bold")).pack(anchor="w")
        tk.Label(izq, text="Llena las plantillas de DBNeT directo desde Workiva",
                 bg=AZUL, fg="#CFE0F0", font=(FUENTE, 10)).pack(anchor="w")
        der = tk.Frame(cab, bg=AZUL)
        der.pack(side="right", anchor="s")
        tk.Label(der, text="Carpeta de trabajo", bg=AZUL, fg="#CFE0F0",
                 font=(FUENTE, 8)).pack(anchor="e")
        fila = tk.Frame(der, bg=AZUL)
        fila.pack(anchor="e")
        tk.Label(fila, textvariable=self.info_carpeta, bg=AZUL, fg=BLANCO,
                 font=(FUENTE, 9)).pack(side="left")
        self.b_cambia_carpeta = self._enlace(fila, "Cambiar", self._cambia_carpeta,
                                             bg=AZUL, fg="#CFE0F0")
        self.b_cambia_carpeta.pack(side="left", padx=(8, 0))

        pie = tk.Frame(self, bg=FONDO, padx=24)
        pie.pack(side="bottom", fill="x")
        self.b_log = self._enlace(pie, "detalle técnico", self._muestra_log,
                                  bg=FONDO, fg=TENUE)
        self.b_log.config(font=(FUENTE, 8, "underline"))
        self.b_log.pack(side="right", pady=(0, 4))
        cuerpo = tk.Frame(self, bg=FONDO, padx=24, pady=14)
        cuerpo.pack(fill="both", expand=True)
        cuerpo.columnconfigure(0, weight=1)
        cuerpo.rowconfigure(1, weight=1)

        # ---- que llenar, el boton y el avance, en un solo bloque: separados
        # le quitaban al resultado el espacio para mostrar los avisos.
        tarjeta = self._tarjeta(cuerpo, 0)
        fila = tk.Frame(tarjeta, bg=BLANCO)
        fila.pack(fill="x")
        self.b_llenar = tk.Button(fila, text="Llenar", command=self.llenar,
                                  bg=VERDE, fg=BLANCO, activebackground=VERDE,
                                  activeforeground=BLANCO, relief="flat", bd=0,
                                  padx=34, pady=7, cursor="hand2",
                                  font=(FUENTE, 11, "bold"))
        self.b_llenar.pack(side="right")
        self._etiqueta(fila, "Empresa").pack(side="left")
        self.c_empresa = ttk.Combobox(fila, textvariable=self.empresa, width=10,
                                      state="readonly", font=(FUENTE, 11))
        self.c_empresa.pack(side="left", padx=(8, 24))
        self.c_empresa.bind("<<ComboboxSelected>>", lambda e: self._cambia_eleccion())
        self._etiqueta(fila, "Período").pack(side="left")
        self.c_periodo = ttk.Combobox(fila, textvariable=self.periodo, width=10,
                                      state="readonly", font=(FUENTE, 11))
        self.c_periodo.pack(side="left", padx=(8, 20))
        self.c_periodo.bind("<<ComboboxSelected>>", lambda e: self._cambia_eleccion())
        self.b_buscar = self._enlace(fila, "↻  Buscar en Workiva", self._busca_planillas)
        self.b_buscar.pack(side="left")

        self.l_workiva = tk.Label(tarjeta, textvariable=self.info_workiva,
                                  bg=BLANCO, fg=GRIS, font=(FUENTE, 9), anchor="w")
        self.l_workiva.pack(fill="x", pady=(10, 0))
        fila_pl = tk.Frame(tarjeta, bg=BLANCO)
        fila_pl.pack(fill="x", pady=(2, 0))
        self.l_plantillas = tk.Label(fila_pl, textvariable=self.info_plantillas,
                                     bg=BLANCO, fg=GRIS, font=(FUENTE, 9), anchor="w")
        self.l_plantillas.pack(side="left")
        self.b_cargar = self._enlace(fila_pl, "Cargar plantillas de DBNeT…",
                                     self._carga_plantillas)
        self.b_ver_plantillas = self._enlace(fila_pl, "Ver plantillas",
                                             self._abre_plantillas)

        # El avance solo aparece mientras trabaja: quieto no dice nada.
        self.avance = tk.Frame(tarjeta, bg=BLANCO)
        tk.Frame(self.avance, bg="#E6EAF0", height=1).pack(fill="x", pady=(12, 10))
        tk.Label(self.avance, textvariable=self.estado, bg=BLANCO, fg="#222",
                 font=(FUENTE, 10), anchor="w").pack(fill="x")
        self.barra = ttk.Progressbar(self.avance, mode="determinate", maximum=100)
        self.barra.pack(fill="x", pady=(6, 0))

        # ---- resultado: se arma de nuevo cada vez que cambia
        self.t_resultado = self._tarjeta(cuerpo, 1, "Resultado", estira=True)
        self.resultado = tk.Frame(self.t_resultado, bg=BLANCO)
        self.resultado.pack(fill="both", expand=True)
        self.resultado.bind("<Configure>", self._ajusta_textos)

    def _tarjeta(self, padre, fila, titulo=None, estira=False):
        marco = tk.Frame(padre, bg=BLANCO, padx=18, pady=12,
                         highlightbackground="#DDE3EA", highlightthickness=1)
        marco.grid(row=fila, column=0, sticky="nsew" if estira else "ew",
                   pady=(0, 12))
        if titulo:
            tk.Label(marco, text=titulo.upper(), bg=BLANCO, fg=TENUE,
                     font=(FUENTE, 8, "bold")).pack(anchor="w", pady=(0, 6))
        return marco

    def _etiqueta(self, padre, texto):
        return tk.Label(padre, text=texto, bg=BLANCO, fg="#222",
                        font=(FUENTE, 10, "bold"))

    def _enlace(self, padre, texto, accion, bg=BLANCO, fg=AZUL):
        return tk.Button(padre, text=texto, command=accion, bg=bg, fg=fg,
                         activebackground=bg, activeforeground=fg,
                         relief="flat", bd=0, cursor="hand2", padx=2,
                         font=(FUENTE, 9, "underline"))

    def _boton(self, padre, texto, accion, principal=False):
        color = AZUL if principal else BLANCO
        b = tk.Button(padre, text=texto, command=accion,
                      bg=color, fg=BLANCO if principal else AZUL_OSC,
                      activebackground=AZUL_OSC if principal else "#E8EEF5",
                      activeforeground=BLANCO if principal else AZUL_OSC,
                      relief="flat" if principal else "solid", bd=0 if principal else 1,
                      padx=16, pady=6, cursor="hand2",
                      font=(FUENTE, 9, "bold" if principal else "normal"))
        return b

    # ----------------------------------------------------- carpeta de trabajo
    def _refresca_carpeta(self):
        texto = str(BASE)
        if len(texto) > 60:
            texto = "…" + texto[-58:]
        self.info_carpeta.set(texto)

    def _cambia_carpeta(self):
        global BASE
        if self.corriendo:
            return
        elegida = filedialog.askdirectory(
            parent=self, initialdir=str(BASE),
            title="Carpeta de trabajo: donde van (o están) las carpetas de cada empresa")
        if not elegida:
            return
        BASE = Path(elegida)
        guarda_config({**lee_config(), "carpeta": str(BASE)})
        self._refresca_carpeta()
        self._refresca_empresas()
        self._cambia_eleccion()

    # -------------------------------------------------- empresas y periodos
    def _lee_cache(self):
        try:
            return [Planilla(**d) for d in json.loads(CACHE.read_text("utf-8"))]
        except (OSError, ValueError, TypeError):
            return []

    def _guarda_cache(self):
        try:
            DATOS.mkdir(parents=True, exist_ok=True)
            CACHE.write_text(json.dumps([p.__dict__ for p in self.planillas]), "utf-8")
        except OSError:
            pass

    def _refresca_empresas(self):
        todas = sorted(set(empresas_locales()) | {p.empresa for p in self.planillas})
        self.c_empresa["values"] = todas
        if self.empresa.get() not in todas:
            # Por omision la ultima que se uso, o la primera con plantillas.
            ultima = lee_config().get("empresa")
            con_plantillas = [e for e in todas if plantillas_de(BASE / e)]
            self.empresa.set(ultima if ultima in todas
                             else (con_plantillas or todas or [""])[0])
        self._cambia_eleccion()

    def _refresca_periodos(self):
        propias = [p for p in self.planillas if p.empresa == self.empresa.get()]
        propias.sort(key=lambda p: (p.orden, p.modificada), reverse=True)
        valores = list(dict.fromkeys(p.periodo for p in propias))
        self.c_periodo["values"] = valores
        if self.periodo.get() not in valores:
            # El ultimo que se eligio para esta empresa; si no, el mas nuevo.
            recordado = lee_config().get("periodos", {}).get(self.empresa.get())
            self.periodo.set(recordado if recordado in valores
                             else (valores[0] if valores else ""))
        self._refresca_info()

    def _cambia_eleccion(self):
        """Otra empresa o periodo: se muestra lo que hay en disco para eso."""
        self._refresca_periodos()
        if self.corriendo:
            return
        self.avance.pack_forget()
        if self.empresa.get():
            config = lee_config()
            periodos = {**config.get("periodos", {})}
            if self.periodo.get():
                periodos[self.empresa.get()] = self.periodo.get()
            guarda_config({**config, "empresa": self.empresa.get(),
                           "periodos": periodos})
        self._muestra_guardado()

    def _planilla(self):
        return next((p for p in self.planillas if p.empresa == self.empresa.get()
                     and p.periodo == self.periodo.get()), None)

    def _refresca_info(self):
        emp = self.empresa.get()
        p = self._planilla()
        if p:
            f = p.modificada_local()
            cuando = f" · modificada el {f:%d-%m-%Y a las %H:%M}" if f else ""
            iguales = sum(1 for q in self.planillas
                          if q.empresa == p.empresa and q.periodo == p.periodo)
            extra = (f"   (hay {iguales} planillas con este nombre; se usa la "
                     "modificada más recientemente)" if iguales > 1 else "")
            self.info_workiva.set(f"En Workiva: «{p.nombre}»{cuando}{extra}")
            self.l_workiva.config(fg=AMBAR if iguales > 1 else GRIS)
        elif self.buscando:
            self.info_workiva.set("Buscando las planillas XBRL en Workiva…")
            self.l_workiva.config(fg=GRIS)
        elif emp:
            self.info_workiva.set(f"En Workiva no hay ninguna planilla «{emp} XBRL MM-AAAA».")
            self.l_workiva.config(fg=ROJO)
        else:
            self.info_workiva.set("")

        self.b_cargar.pack_forget()
        self.b_ver_plantillas.pack_forget()
        if not emp:
            self.info_plantillas.set("")
            return
        n = len(plantillas_de(BASE / emp))
        if n:
            self.info_plantillas.set(f"Plantillas de DBNeT de {emp}: {n} cargadas")
            self.l_plantillas.config(fg=GRIS)
            self.b_cargar.config(text="Cambiar por otras…")
        else:
            self.info_plantillas.set(f"Plantillas de DBNeT de {emp}: todavía no se cargan")
            self.l_plantillas.config(fg=ROJO)
            self.b_cargar.config(text="Cargar plantillas de DBNeT…")
        self.b_cargar.pack(side="left", padx=(10, 0))
        if n:
            self.b_ver_plantillas.pack(side="left", padx=(10, 0))

    def _carga_plantillas(self):
        emp = self.empresa.get()
        if not emp or self.corriendo:
            return
        origen = filedialog.askdirectory(
            parent=self, mustexist=True,
            title=f"Elige la CARPETA con las plantillas de DBNeT de {emp} "
                  "(.xlsm sueltos o el .zip de DBNeT)")
        if not origen:
            return
        self.config(cursor="watch")
        self.update_idletasks()
        try:
            nuevas, descartes, repetidas, propias = reune_plantillas(origen)
        finally:
            self.config(cursor="")
        if repetidas:
            lista = "\n".join(f"  • {n}\n      en: " + "\n      y en: ".join(r[:3])
                              for n, r in list(repetidas.items())[:5])
            if len(repetidas) > 5:
                lista += f"\n  … y {len(repetidas) - 5} más"
            messagebox.showerror(
                "Plantillas repetidas",
                f"En esa carpeta hay {plural(len(repetidas), 'plantilla', 'plantillas')} "
                "con el mismo nombre en más de un lugar, y no son iguales (una "
                "puede estar ya llenada):\n\n" + lista +
                "\n\nNo se cargó nada. Elige la carpeta que tiene solo las "
                "plantillas que entregó DBNeT.", parent=self)
            return
        motivos = "\n".join(f"  • {n}: {m}" for n, m in descartes[:12])
        if len(descartes) > 12:
            motivos += f"\n  … y {len(descartes) - 12} más"
        if not nuevas:
            texto = f"En {origen} no hay plantillas .xlsm de DBNeT."
            if propias:
                texto = ("Esa carpeta es de la propia app: ahí están las plantillas "
                         "ya cargadas (o llenadas). Elige la carpeta donde dejaste "
                         "las plantillas que entregó DBNeT.")
            if motivos:
                texto += f"\n\nSe descartaron:\n{motivos}"
            messagebox.showerror("No hay plantillas", texto, parent=self)
            return

        carpeta = BASE / emp
        actuales = len(plantillas_de(carpeta))
        texto = f"Se encontraron {len(nuevas)} plantillas de DBNeT para {emp}."
        if actuales:
            texto += (f"\n\nReemplazan a las {actuales} que tiene ahora. Esas no se "
                      "borran: quedan guardadas aparte, en la carpeta "
                      f"{emp}\\plantillas_anteriores.")
        if descartes:
            texto += f"\n\nNo se van a cargar:\n{motivos}"
        if not messagebox.askyesno("Cargar plantillas", texto + "\n\n¿Cargarlas?",
                                   parent=self):
            return
        try:
            instala_plantillas(carpeta, nuevas)
        except OSError as e:
            messagebox.showerror("No se pudieron cargar", str(e), parent=self)
            return
        self._cambia_eleccion()
        messagebox.showinfo("Plantillas cargadas",
                            f"Listo: {emp} tiene {len(nuevas)} plantillas de DBNeT cargadas.",
                            parent=self)

    def _abre_plantillas(self):
        emp = self.empresa.get()
        if emp and (BASE / emp / "xls").is_dir():
            abre(BASE / emp / "xls")

    def _busca_planillas(self):
        if self.buscando:
            return
        self.buscando = True
        self.b_buscar.config(state="disabled", text="Buscando en Workiva…")
        self._refresca_info()

        def trabajo():
            try:
                self.cola.put(("planillas", Workiva().planillas()))
            except ErrorWorkiva as e:
                self.cola.put(("planillas_error", str(e)))
            except Exception as e:                       # noqa: BLE001
                self.cola.put(("planillas_error", f"{type(e).__name__}: {e}"))

        threading.Thread(target=trabajo, daemon=True).start()

    # ------------------------------------------------------------ resultado
    def _limpia_resultado(self):
        for w in self.resultado.winfo_children():
            w.destroy()

    def _ajusta_textos(self, evento=None):
        """Los textos se cortan al ancho que tenga la ventana, no a uno fijo."""
        ancho = max(200, self.resultado.winfo_width() - 10)
        for w in self.resultado.winfo_children():
            if isinstance(w, tk.Label):
                w.config(wraplength=ancho)

    def _texto(self, texto, color="#222", peso="normal", tam=10, **pack):
        l = tk.Label(self.resultado, text=texto, bg=BLANCO, fg=color, justify="left",
                     anchor="w", font=(FUENTE, tam, peso),
                     wraplength=max(200, self.resultado.winfo_width() - 10))
        l.pack(fill="x", **({"pady": (0, 4)} | pack))
        return l

    def _muestra_guardado(self):
        """Lo que quedo en disco del ultimo llenado de esta empresa."""
        self._limpia_resultado()
        emp, p = self.empresa.get(), self._planilla()
        if not emp:
            return
        ultimo = lee_ultimo(BASE / emp)
        if ultimo and p and ultimo.get("planilla") == p.nombre:
            archivo = BASE / emp / ultimo.get("archivo", "")
            if archivo.is_file():
                return self._muestra_listo(emp, ultimo, archivo, guardado=True)
        nombre = p.nombre if p else f"{emp} {self.periodo.get()}".strip()
        # Se lleno antes, y despues se lleno otro periodo encima: las
        # plantillas ya no son de este periodo, pero su archivo para DBNeT
        # sigue siendo valido, porque trae sus propias copias de las hojas.
        archivo = next((a for a in (BASE / emp / f"{nombre}_LLENADO.xlsm",
                                    BASE / emp / f"{nombre}_LLENADO.xlsx")
                        if a.is_file()), None)
        if archivo:
            cuando = time.strftime("%d-%m-%Y a las %H:%M",
                                   time.localtime(archivo.stat().st_mtime))
            self._texto(f"✔  El archivo para DBNeT de «{nombre}» sigue en la "
                        f"carpeta (armado el {cuando}) y se puede usar.",
                        VERDE, "bold", 11)
            caja = tk.Frame(self.resultado, bg="#F2F7FC", padx=14, pady=10,
                            highlightbackground="#D5E3F1", highlightthickness=1)
            caja.pack(fill="x", pady=(6, 10))
            self._enlace(caja, "Mostrarlo en su carpeta",
                         lambda: muestra_en_carpeta(archivo), bg="#F2F7FC").pack(
                             side="right", padx=(10, 0))
            self._boton(caja, "Abrir archivo", lambda: abre(archivo),
                        principal=True).pack(side="right")
            tk.Label(caja, text=f"Archivo para DBNeT:  {archivo.name}", bg="#F2F7FC",
                     fg="#222", font=(FUENTE, 10, "bold"), anchor="w").pack(
                         side="left", fill="x")
            if ultimo:
                self._texto(f"Las plantillas de {emp} tienen hoy los datos de "
                            f"«{ultimo.get('planilla')}». Si vuelves a llenar "
                            "este período, se llenan de nuevo y el archivo se "
                            "reemplaza.", TENUE, tam=9)
            return
        self._texto(f"Todavía no se ha llenado «{nombre}».", GRIS, tam=10)
        if ultimo:
            self._texto(f"Las plantillas de {emp} tienen hoy los datos de "
                        f"«{ultimo.get('planilla')}» (llenadas el "
                        f"{ultimo.get('cuando', '?')}). Al llenar este período se "
                        "reemplazan.", TENUE, tam=9)

    def _muestra_listo(self, emp, ultimo, archivo, guardado=False):
        self._limpia_resultado()
        celdas, avisos = ultimo.get("celdas"), ultimo.get("avisos", 0)
        n = ultimo.get("plantillas")
        titulo = f"✔  «{ultimo['planilla']}» está llenado"
        if celdas is not None:
            titulo += f": {miles(celdas)} celdas"
            if n:
                titulo += f" en {n} plantillas"
        if guardado:
            titulo += f"  ·  el {ultimo.get('cuando', '?')}"
        self._texto(titulo, VERDE, "bold", 11)

        # El archivo que sirve para DBNeT, con lo que hay que hacer con el
        caja = tk.Frame(self.resultado, bg="#F2F7FC", padx=14, pady=10,
                        highlightbackground="#D5E3F1", highlightthickness=1)
        caja.pack(fill="x", pady=(6, 10))
        linea = tk.Frame(caja, bg="#F2F7FC")
        linea.pack(fill="x")
        # Los botones se empacan antes que el nombre: si la ventana es angosta
        # se recorta el nombre, no los botones.
        self._enlace(linea, "Mostrarlo en su carpeta", lambda: muestra_en_carpeta(archivo),
                     bg="#F2F7FC").pack(side="right", padx=(10, 0))
        self._boton(linea, "Abrir archivo", lambda: abre(archivo),
                    principal=True).pack(side="right")
        tk.Label(linea, text=f"Archivo para DBNeT:  {archivo.name}", bg="#F2F7FC",
                 fg="#222", font=(FUENTE, 10, "bold"), anchor="w").pack(
                     side="left", fill="x")
        if archivo.suffix.lower() == ".xlsm":
            ayuda = ("Ábrelo y usa su botón para generar los CSV que se suben a "
                     "DBNeT. Trae todos los cuadros juntos.")
        else:
            ayuda = ("Sin macros: en este computador no se pudo usar Excel para "
                     f"armarlo. Los CSV se pueden generar desde cada plantilla de "
                     f"{emp}, que sí quedaron llenas y con sus botones.")
        ayuda_l = tk.Label(caja, text=ayuda, bg="#F2F7FC", fg=GRIS, font=(FUENTE, 9),
                           anchor="w", justify="left")
        ayuda_l.pack(fill="x", pady=(4, 0))
        caja.bind("<Configure>", lambda e: ayuda_l.config(
            wraplength=max(200, e.width - 40)))

        # Lo que hay que revisar antes de entregar
        revisar = BASE / emp / "REVISAR.xlsx"
        if not avisos or not revisar.exists():
            self._texto("✔  Nada que revisar: todo calzó sin suponer nada.", VERDE)
            return
        fila = tk.Frame(self.resultado, bg=BLANCO)
        fila.pack(fill="x", pady=(2, 4))
        tk.Label(fila, text=f"⚠  Antes de entregar, revisa "
                            f"{plural(avisos, 'aviso', 'avisos')}:",
                 bg=BLANCO, fg=AMBAR, font=(FUENTE, 10, "bold")).pack(side="left")
        self._enlace(fila, "Abrir REVISAR.xlsx", lambda: abre(revisar)).pack(
            side="left", padx=(10, 0))
        self._lista_avisos(revisar)

    def _lista_avisos(self, ruta):
        """Cada aviso completo, uno bajo otro, con barra si no caben. En una
        tabla los textos quedaban cortados y habia que elegir fila por fila."""
        marco = tk.Frame(self.resultado, bg=BLANCO)
        marco.pack(fill="both", expand=True)
        texto = tk.Text(marco, bg=AMBAR_FONDO, fg="#222", bd=0, wrap="word",
                        padx=14, pady=10, font=(FUENTE, 9), height=4,
                        spacing1=1, spacing3=1, cursor="arrow")
        barra = ttk.Scrollbar(marco, orient="vertical", command=texto.yview)
        texto.configure(yscrollcommand=barra.set)
        barra.pack(side="right", fill="y")
        texto.pack(side="left", fill="both", expand=True)
        texto.tag_config("cuadro", font=(FUENTE, 10, "bold"), foreground=AMBAR,
                         spacing1=6)
        texto.tag_config("titulo", font=(FUENTE, 9, "bold"))
        try:
            from openpyxl import load_workbook
            libro = load_workbook(ruta, read_only=True)
            n = 0
            for i, fila in enumerate(libro.active.iter_rows(values_only=True)):
                if not i or not any(fila):
                    continue
                hoja, donde, paso, mirar = ([v if v is not None else ""
                                             for v in fila] + [""] * 4)[:4]
                n += 1
                if n > 1:
                    texto.insert("end", "\n")
                texto.insert("end", f"{n}.  {hoja}\n", "cuadro")
                for titulo, valor in (("Dónde mirar", donde), ("Qué pasó", paso),
                                      ("Qué hay que revisar", mirar)):
                    if valor:
                        texto.insert("end", f"{titulo}: ", "titulo")
                        texto.insert("end", f"{valor}\n")
            libro.close()
        except Exception as e:                           # noqa: BLE001
            texto.insert("end", f"No se pudo leer {ruta.name}: {e}")
        texto.config(state="disabled")

    def _muestra_error(self, paso, texto):
        self._limpia_resultado()
        self._texto(f"✖  No se pudo {FALLO[paso]}", ROJO, "bold", 11)
        self._texto(texto, "#222", tam=9)
        if "a medias" not in texto:
            self._texto("Nada quedó a medias: al volver a llenar se parte otra "
                        "vez de las plantillas originales de DBNeT.", TENUE, tam=9)

    # ------------------------------------------------------------ llenado
    def llenar(self):
        if self.corriendo:
            return
        emp, p = self.empresa.get(), self._planilla()
        if not emp or not p:
            messagebox.showwarning("Falta elegir",
                                   "Elige una empresa y un período que existan en Workiva.",
                                   parent=self)
            return
        if not plantillas_de(BASE / emp):
            messagebox.showwarning(
                "Faltan las plantillas",
                f"{emp} todavía no tiene sus plantillas de DBNeT.\n\n"
                "Usa «Cargar plantillas de DBNeT…» y elige la carpeta donde "
                "están las que entregó DBNeT para esta empresa.", parent=self)
            return

        self.corriendo = True
        for b in (self.b_llenar, self.b_buscar, self.b_cargar, self.b_cambia_carpeta):
            b.config(state="disabled")
        self.c_empresa.config(state="disabled")
        self.c_periodo.config(state="disabled")
        self._limpia_resultado()
        self._texto("Trabajando… el resultado aparece aquí al terminar.", TENUE)
        self.registro = []
        if self.ventana_log:
            self.log.config(state="normal")
            self.log.delete("1.0", "end")
            self.log.config(state="disabled")
        self.avance.pack(fill="x")
        self.barra.config(mode="determinate")
        self.barra["value"] = 0
        self.confirmacion = queue.Queue()
        threading.Thread(target=self._trabaja, args=(emp, p), daemon=True).start()

    def _trabaja(self, emp, planilla):
        """En un hilo aparte; habla con la ventana solo por la cola."""
        carpeta = BASE / emp
        xls = carpeta / "xls"
        revisar = carpeta / "REVISAR.xlsx"
        temporal = TEMPORAL / emp
        # Lo que baja de Workiva es solo material de trabajo: va a la carpeta
        # temporal y no a la de la empresa, donde parecia un export que habia
        # que guardar. Se reemplaza en cada llenado.
        export = temporal / f"{planilla.nombre}.xlsx"
        temporal.mkdir(parents=True, exist_ok=True)
        base = carpeta / f"{planilla.nombre}_LLENADO"
        comun = ["--plantillas", str(xls), "--workiva", str(export),
                 "--sobre-plantillas", "--revisar", str(revisar),
                 "--reporte", str(temporal / "reporte_llenado.csv")]
        paso = 0
        try:
            # 1. Workiva
            self.cola.put(("paso", paso))
            t0 = time.time()
            Workiva().descarga(planilla, export)
            self.cola.put(("linea", (f"Descargado de Workiva en {time.time() - t0:.0f} s: "
                                     f"{export}", "ok")))

            # 2. Simulacion
            paso = 1
            self.cola.put(("paso", paso))
            codigo, lineas = self._corre("llenar", comun + ["--dry-run"])
            if codigo:
                return self._falla(paso, lineas)
            celdas, avisos = self._resumen(lineas)
            n = len(plantillas_de(carpeta))
            self.cola.put(("confirma", (emp, planilla.nombre, celdas, avisos, n)))
            if not self.confirmacion.get():
                return self.cola.put(("fin", ("cancelado", paso, None)))

            # 3. Llenado
            paso = 2
            self.cola.put(("paso", paso))
            # Desde aqui las plantillas cambian: lo que decia el ultimo
            # llenado deja de ser cierto hasta que este termine.
            nota_ultimo(carpeta).unlink(missing_ok=True)
            codigo, lineas = self._corre("llenar", comun)
            if codigo:
                return self._falla(paso, lineas)
            celdas, avisos = self._resumen(lineas)

            # 4. Archivo unico
            paso = 3
            self.cola.put(("paso", paso))
            # Igual que el .bat: si el .xlsm con macros no sale (sin Excel en
            # el equipo), queda al menos el .xlsx. Un archivo abierto, en
            # cambio, es para cerrarlo y volver a correr, no para seguir.
            codigo_macros, lineas = self._corre(
                "fusionar", ["--origen", str(xls), "--salida", f"{base}.xlsm",
                             "--con-macros", "--solo-workiva"])
            # Solo la falta de Excel/pywin32 tiene salida por el .xlsx; cualquier
            # otro corte (un archivo abierto, o que el control de datos no
            # calzo) se muestra, no se tapa con un archivo sin macros.
            if codigo_macros and not any("hace falta pywin32" in l for l in lineas):
                return self._falla(paso, lineas)
            # El .xlsx sin macros es solo el plan B: si el .xlsm salio, sobra
            # (y dos archivos con el mismo nombre confundian cual entregar).
            if codigo_macros:
                codigo, lineas_x = self._corre(
                    "fusionar", ["--origen", str(xls), "--salida", f"{base}.xlsx",
                                 "--solo-workiva"])
                if codigo:
                    return self._falla(paso, lineas_x)
                final = Path(f"{base}.xlsx")
            else:
                final = Path(f"{base}.xlsm")
                Path(f"{base}.xlsx").unlink(missing_ok=True)
            ultimo = {"planilla": planilla.nombre, "archivo": final.name,
                      "celdas": celdas, "avisos": avisos, "plantillas": n,
                      "cuando": time.strftime("%d-%m-%Y a las %H:%M")}
            nota_ultimo(carpeta).parent.mkdir(parents=True, exist_ok=True)
            nota_ultimo(carpeta).write_text(json.dumps(ultimo, indent=1), "utf-8")
            self.cola.put(("fin", ("ok", paso, (emp, ultimo, final))))
        except ErrorWorkiva as e:
            self.cola.put(("linea", (str(e), "err")))
            self.cola.put(("fin", ("error", paso, str(e))))
        except Exception as e:                           # noqa: BLE001
            self.cola.put(("linea", (traceback.format_exc(), "err")))
            self.cola.put(("fin", ("error", paso, f"{type(e).__name__}: {e}")))

    def _corre(self, tarea, args):
        """Corre una tarea como proceso aparte y devuelve (codigo, lineas)."""
        if getattr(sys, "frozen", False):
            orden = [sys.executable, "--tarea", tarea, *args]
        else:
            orden = [sys.executable, "-u", str(Path(__file__).resolve()),
                     "--tarea", tarea, *args]
        respaldo = Path(tempfile.gettempdir()) / f"xbrl_{tarea}_{os.getpid()}.log"
        respaldo.unlink(missing_ok=True)
        entorno = {**os.environ, "PYTHONIOENCODING": "utf-8",
                   "PYTHONUNBUFFERED": "1", "XBRL_SALIDA": str(respaldo)}
        self.cola.put(("linea", (f"\n── {tarea} " + "─" * 50, "paso")))
        proc = subprocess.Popen(orden, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                stdin=subprocess.DEVNULL, text=True, encoding="utf-8",
                                errors="replace", cwd=str(BASE), env=entorno,
                                creationflags=SIN_CONSOLA)
        lineas = []
        for linea in proc.stdout:
            lineas.append(linea.rstrip())
            self.cola.put(("linea", (lineas[-1], None)))
        codigo = proc.wait()
        if not lineas and respaldo.exists():
            for linea in respaldo.read_text("utf-8", "replace").splitlines():
                lineas.append(linea)
                self.cola.put(("linea", (linea, None)))
        respaldo.unlink(missing_ok=True)
        return codigo, lineas

    @staticmethod
    def _resumen(lineas):
        celdas, avisos = None, 0
        for l in lineas:
            m = TOTAL.match(l.strip())
            if m:
                celdas = int(re.sub(r"\D", "", m.group(1)))
            m = AVISOS.search(l)
            if m:
                avisos = int(m.group(1))
        return celdas, avisos

    def _falla(self, paso, lineas):
        """El mensaje para la persona: lo que el script explico al cortar.

        Los sys.exit con texto dicen que hacer; un traceback no le dice nada
        a nadie, asi que de el queda solo la ultima linea y el resto va al
        detalle tecnico."""
        if any(l.startswith("Traceback") for l in lineas):
            ultima = next((l.strip() for l in reversed(lineas) if l.strip()), "")
            texto = (f"{ultima}\n\nEsto es un error del programa, no tuyo. "
                     "Manda una captura de «detalle técnico» para revisarlo.")
        else:
            utiles = [l.rstrip() for l in lineas
                      if l.strip() and not AVANCE.match(l)]
            # Desde donde empieza el ultimo mensaje (la linea sin sangria
            # mas cercana al final, o las ultimas 12).
            inicio = max(0, len(utiles) - 12)
            for i in range(len(utiles) - 1, inicio - 1, -1):
                if utiles[i].startswith("  ") and not utiles[i].startswith("   "):
                    if i == 0 or not utiles[i - 1].startswith("  "):
                        inicio = i
                        break
            # Las lineas con mas sangria continuan la anterior: el script
            # corta a 70 columnas para la consola, y aqui eso sobra.
            parrafos = []
            for l in utiles[inicio:]:
                if parrafos and l.startswith("     "):
                    parrafos[-1] += " " + l.strip()
                else:
                    parrafos.append(l.strip())
            texto = "\n".join(parrafos)
        texto = texto or "El proceso terminó sin decir por qué."
        self.cola.put(("fin", ("error", paso, texto)))

    # ------------------------------------------------------------ recepcion
    def _drena(self):
        try:
            while True:
                tipo, dato = self.cola.get_nowait()
                getattr(self, f"_en_{tipo}")(dato)
        except queue.Empty:
            pass
        self.after(80, self._drena)

    def _en_planillas(self, lista):
        self.buscando = False
        self.planillas = lista
        self._guarda_cache()
        self.b_buscar.config(state="normal" if not self.corriendo else "disabled",
                             text="↻  Buscar en Workiva")
        if self.corriendo:
            self._refresca_info()
        else:
            self._refresca_empresas()

    def _en_planillas_error(self, texto):
        self.buscando = False
        self.b_buscar.config(state="normal" if not self.corriendo else "disabled",
                             text="↻  Buscar en Workiva")
        self._refresca_info()
        if not self.planillas:
            messagebox.showerror("Workiva", texto, parent=self)
        else:
            self.info_workiva.set(self.info_workiva.get() + "   (no se pudo "
                                  "actualizar la lista de Workiva; es la última "
                                  "que se encontró)")

    def _en_linea(self, dato):
        texto, tag = dato
        self._anota(texto, tag)
        m = N_PLANTILLAS.match(texto)
        if m:
            self.total, self.hechos = int(m.group(1)), 0
        elif self.total and AVANCE.match(texto):
            self.hechos += 1
            self.barra["value"] = min(100, self.hechos * 100 / self.total)
            self.estado.set(f"{PASOS[self.paso]}  ({self.hechos} de {self.total})")

    def _en_paso(self, paso):
        self.paso = paso
        self.estado.set(PASOS[paso])
        self.total = self.hechos = 0
        if paso in (0, 3):
            self.barra.config(mode="indeterminate")
            self.barra.start(12)
        else:
            self.barra.stop()
            self.barra.config(mode="determinate")
            self.barra["value"] = 0

    def _en_confirma(self, dato):
        emp, nombre, celdas, avisos, n = dato
        texto = (f"Los datos de «{nombre}» están listos para pasar a las "
                 f"{n} plantillas de {emp}"
                 + (f" ({miles(celdas)} celdas)" if celdas is not None else "") + ".")
        if avisos:
            texto += (f"\n\nHay {plural(avisos, 'cosa', 'cosas')} para revisar: "
                      "las vas a ver en «Resultado» al terminar.")
        texto += ("\n\nLas plantillas se sobrescriben; las originales de DBNeT "
                  "quedan respaldadas y se reponen en cada llenado.\n\n¿Continuar?")
        self.barra.stop()
        self.confirmacion.put(messagebox.askyesno("Confirmar llenado", texto,
                                                  parent=self))

    def _en_fin(self, dato):
        tipo, paso, info = dato
        self.corriendo = False
        self.barra.stop()
        self.barra.config(mode="determinate")
        for b in (self.b_llenar, self.b_buscar, self.b_cargar, self.b_cambia_carpeta):
            b.config(state="normal")
        self.c_empresa.config(state="readonly")
        self.c_periodo.config(state="readonly")

        self.avance.pack_forget()
        if tipo == "ok":
            emp, ultimo, final = info
            self._anota(f"\nListo: {final}", "ok")
            self._muestra_listo(emp, ultimo, final)
        elif tipo == "cancelado":
            self._muestra_guardado()
            self._texto("Cancelado: no se escribió ningún archivo.", AMBAR, tam=9)
        else:
            self._anota(info, "err")
            self._muestra_error(paso, info)
            self.bell()

    # ------------------------------------------------------------ utilidades
    def _anota(self, texto, tag=None):
        self.registro.append((texto, tag))
        if self.ventana_log:
            self._escribe_log(texto, tag)

    def _escribe_log(self, texto, tag):
        self.log.config(state="normal")
        self.log.insert("end", texto + "\n", tag or ())
        self.log.see("end")
        self.log.config(state="disabled")

    def _muestra_log(self):
        """Lo que imprimieron los scripts, para diagnosticar. Aparte, porque
        a la vista no le sirve a nadie mas."""
        if self.ventana_log:
            self.ventana_log.deiconify()
            self.ventana_log.lift()
            return
        v = self.ventana_log = tk.Toplevel(self)
        v.title("Enlace XBRL — detalle técnico")
        v.geometry(f"{int(900 * self.escala)}x{int(500 * self.escala)}")
        self.log = tk.Text(v, bg="#1E1E1E", fg="#D4D4D4", bd=0,
                           font=("Consolas", 9), wrap="none", padx=10, pady=8)
        barra = ttk.Scrollbar(v, orient="vertical", command=self.log.yview)
        self.log.configure(yscrollcommand=barra.set)
        self.log.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        self.log.tag_config("err", foreground="#F48771")
        self.log.tag_config("ok", foreground="#7BD88F")
        self.log.tag_config("paso", foreground="#9CDCFE")
        if not self.registro:
            self._escribe_log("Todavía no se ha llenado nada en esta sesión.", None)
        for texto, tag in self.registro:
            self._escribe_log(texto, tag)

        def cierra():
            self.ventana_log.destroy()
            self.ventana_log = None
        v.protocol("WM_DELETE_WINDOW", cierra)


def main():
    if os.name == "nt":
        # Sin esto Windows estira la ventana y el texto sale borroso en
        # pantallas con escala.
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except (AttributeError, OSError):
            pass
    Aplicacion().mainloop()


if __name__ == "__main__":
    main()
