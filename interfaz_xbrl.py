#!/usr/bin/env python3
"""Ventana para el llenado XBRL de DBNeT desde el export de Workiva.

Hace lo mismo que LLENAR_XBRL.bat -- simular, llenar las 41 plantillas y
fusionarlas en un solo archivo -- pero mostrando en que va, y dejando a la
vista las celdas que hay que mirar antes de entregar, que es lo unico que
de verdad pide atencion.

Los dos scripts corren como procesos aparte y su salida se lee linea a
linea: asi la ventana sigue respondiendo y el registro va apareciendo
mientras trabaja, en vez de congelarse hasta el final.
"""

import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

BASE = Path(__file__).resolve().parent
LLENAR = BASE / "llenar_dbnet_desde_workiva.py"
FUSIONAR = BASE / "fusionar_cuadros.py"
TEMPORAL = Path(tempfile.gettempdir()) / "XBRL_llenado"

AZUL, AZUL_OSC, GRIS, BLANCO = "#0B5394", "#083D6E", "#F4F6F8", "#FFFFFF"
VERDE, ROJO = "#1E7B34", "#B3261E"

# "  IAS12-Cuadros(835110)_2025.xlsm    202 celdas  (dry-run)", una por
# plantilla. Se ancla en el campo de ancho fijo del nombre y no en el ".xlsm",
# porque el script recorta los nombres largos y a tres de las 41 les come la
# extension. Ese ancho es ademas lo que distingue esta linea del aviso
# "  OJO: N celdas necesitan que las mires", que tambien dice "celdas".
AVANCE = re.compile(r"^\s{2}.{50}\s*\d+\s+celdas\b")
TOTAL = re.compile(r"^Plantillas:\s*(\d+)")

# Sin consola negra detras de la ventana al lanzar cada script.
SIN_CONSOLA = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0


class Interfaz(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("XBRL DBNeT — Llenado desde Workiva")
        self.configure(bg=GRIS)
        self.geometry("980x680")
        self.minsize(860, 600)

        self.plantillas = tk.StringVar(value=str(BASE / "xls"))
        self.workiva = tk.StringVar()
        self.estado = tk.StringVar(value="Elige el export de Workiva para empezar.")
        self.cola = queue.Queue()
        self.corriendo = False
        self.salida_final = None
        self.hechos = self.total = 0

        self._construye()
        self._busca_workiva()
        self.after(80, self._drena)

    # ------------------------------------------------------------ interfaz
    def _construye(self):
        cab = tk.Frame(self, bg=AZUL, padx=20, pady=14)
        cab.pack(fill="x")
        tk.Label(cab, text="Llenado XBRL", bg=AZUL, fg=BLANCO,
                 font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(cab, text="Plantillas de DBNeT llenadas con los datos de Workiva",
                 bg=AZUL, fg="#CFE0F0", font=("Segoe UI", 10)).pack(anchor="w")

        cuerpo = tk.Frame(self, bg=GRIS, padx=20, pady=16)
        cuerpo.pack(fill="both", expand=True)

        self._selector(cuerpo, "Plantillas de DBNeT", self.plantillas,
                       self._elige_plantillas, 0)
        self._selector(cuerpo, "Export de Workiva", self.workiva,
                       self._elige_workiva, 1)

        botones = tk.Frame(cuerpo, bg=GRIS)
        botones.grid(row=2, column=0, columnspan=3, sticky="w", pady=(14, 8))
        self.b_simular = self._boton(botones, "Simular", self.simular, AZUL_OSC)
        self.b_generar = self._boton(botones, "Generar archivo", self.generar, VERDE)
        self.b_abrir = self._boton(botones, "Abrir resultado", self.abrir, "#5A5A5A")
        self.b_abrir.config(state="disabled")

        self.barra = ttk.Progressbar(cuerpo, mode="determinate", maximum=100)
        self.barra.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(4, 2))
        tk.Label(cuerpo, textvariable=self.estado, bg=GRIS, fg="#333",
                 font=("Segoe UI", 9), anchor="w").grid(
                     row=4, column=0, columnspan=3, sticky="ew")

        pestanas = ttk.Notebook(cuerpo)
        pestanas.grid(row=5, column=0, columnspan=3, sticky="nsew", pady=(12, 0))
        cuerpo.rowconfigure(5, weight=1)
        cuerpo.columnconfigure(1, weight=1)

        marco_rev = tk.Frame(pestanas, bg=BLANCO)
        pestanas.add(marco_rev, text="  Por revisar  ")
        cols = ("hoja", "donde", "paso", "mirar", "casos")
        self.tabla = ttk.Treeview(marco_rev, columns=cols, show="headings")
        for c, txt, ancho in zip(cols,
                                 ("Hoja", "Dónde mirar", "Qué pasó",
                                  "Qué tienes que mirar", "Casos"),
                                 (180, 150, 330, 330, 55)):
            self.tabla.heading(c, text=txt)
            self.tabla.column(c, width=ancho, anchor="w")
        bar_v = ttk.Scrollbar(marco_rev, orient="vertical", command=self.tabla.yview)
        self.tabla.configure(yscrollcommand=bar_v.set)
        self.tabla.pack(side="left", fill="both", expand=True)
        bar_v.pack(side="right", fill="y")

        marco_log = tk.Frame(pestanas, bg=BLANCO)
        pestanas.add(marco_log, text="  Registro  ")
        self.log = tk.Text(marco_log, bg="#1E1E1E", fg="#D4D4D4", bd=0,
                           font=("Consolas", 9), wrap="none", padx=10, pady=8)
        bar_l = ttk.Scrollbar(marco_log, orient="vertical", command=self.log.yview)
        self.log.configure(yscrollcommand=bar_l.set, state="disabled")
        self.log.pack(side="left", fill="both", expand=True)
        bar_l.pack(side="right", fill="y")
        self.log.tag_config("err", foreground="#F48771")
        self.log.tag_config("ok", foreground="#7BD88F")

    def _selector(self, padre, titulo, variable, accion, fila):
        tk.Label(padre, text=titulo, bg=GRIS, fg="#222",
                 font=("Segoe UI", 10, "bold")).grid(
                     row=fila, column=0, sticky="w", pady=4)
        tk.Entry(padre, textvariable=variable, font=("Segoe UI", 9),
                 bg=BLANCO, relief="solid", bd=1).grid(
                     row=fila, column=1, sticky="ew", padx=10, pady=4, ipady=4)
        tk.Button(padre, text="Elegir…", command=accion, bg=BLANCO, bd=1,
                  relief="solid", cursor="hand2", padx=12, pady=3,
                  font=("Segoe UI", 9)).grid(row=fila, column=2, pady=4)

    def _boton(self, padre, texto, accion, color):
        b = tk.Button(padre, text=texto, command=accion, bg=color, fg=BLANCO,
                      activebackground=color, activeforeground=BLANCO,
                      relief="flat", bd=0, padx=20, pady=9, cursor="hand2",
                      font=("Segoe UI", 10, "bold"))
        b.pack(side="left", padx=(0, 10))
        return b

    # -------------------------------------------------------------- eleccion
    def _busca_workiva(self):
        """El export mas reciente de workiva\\, igual que hace el .bat."""
        carpeta = BASE / "workiva"
        if not carpeta.is_dir():
            return
        xlsx = sorted(carpeta.glob("*.xlsx"), key=lambda p: p.stat().st_mtime,
                      reverse=True)
        if xlsx:
            self.workiva.set(str(xlsx[0]))
            self.estado.set(f"Export encontrado: {xlsx[0].name}")

    def _elige_plantillas(self):
        d = filedialog.askdirectory(title="Carpeta con los .xlsm de DBNeT",
                                    initialdir=self.plantillas.get() or str(BASE))
        if d:
            self.plantillas.set(d)

    def _elige_workiva(self):
        f = filedialog.askopenfilename(title="Export de Workiva",
                                       filetypes=[("Excel", "*.xlsx")],
                                       initialdir=str(BASE / "workiva"))
        if f:
            self.workiva.set(f)

    # ------------------------------------------------------------- ejecucion
    def simular(self):
        self._arranca(dry=True)

    def generar(self):
        self._arranca(dry=False)

    def _arranca(self, dry):
        if self.corriendo:
            return
        plantillas, workiva = Path(self.plantillas.get()), Path(self.workiva.get())
        if not plantillas.is_dir() or not list(plantillas.rglob("*.xlsm")):
            messagebox.showerror("Faltan las plantillas",
                                 "En esa carpeta no hay archivos .xlsm de DBNeT.")
            return
        if not workiva.is_file():
            messagebox.showerror("Falta el export",
                                 "Elige el .xlsx que exportaste de Workiva.")
            return

        self.corriendo = True
        self.salida_final = None
        for b in (self.b_simular, self.b_generar, self.b_abrir):
            b.config(state="disabled")
        self.tabla.delete(*self.tabla.get_children())
        self.log.config(state="normal")
        self.log.delete("1.0", "end")
        self.log.config(state="disabled")
        self.hechos, self.total = 0, 0
        self.barra["value"] = 0
        threading.Thread(target=self._trabaja, args=(dry, plantillas, workiva),
                         daemon=True).start()

    def _trabaja(self, dry, plantillas, workiva):
        """Corre en un hilo aparte; se comunica por la cola."""
        try:
            if TEMPORAL.exists():
                shutil.rmtree(TEMPORAL, ignore_errors=True)
            revisar = BASE / "REVISAR.xlsx"
            orden = [sys.executable, "-u", str(LLENAR),
                     "--plantillas", str(plantillas), "--workiva", str(workiva),
                     "--salida", str(TEMPORAL),
                     "--reporte", str(TEMPORAL / "reporte_llenado.csv"),
                     "--revisar", str(revisar)]
            if dry:
                orden.append("--dry-run")
            if self._corre(orden) != 0:
                return self.cola.put(("fin", "El llenado termino con errores."))
            if dry:
                self.cola.put(("revisar", str(revisar)))
                return self.cola.put(("fin", "Simulacion lista. No se escribio nada."))

            base = workiva.stem + "_LLENADO"
            self.cola.put(("paso", "Fusionando en un solo archivo…"))
            con_macros = self._corre(
                [sys.executable, "-u", str(FUSIONAR), "--origen", str(TEMPORAL),
                 "--salida", str(BASE / f"{base}.xlsm"),
                 "--con-macros", "--solo-workiva"]) == 0
            if self._corre(
                [sys.executable, "-u", str(FUSIONAR), "--origen", str(TEMPORAL),
                 "--salida", str(BASE / f"{base}.xlsx"), "--solo-workiva"]) != 0:
                return self.cola.put(("fin", "La fusion termino con errores."))

            self.cola.put(("revisar", str(revisar)))
            if con_macros:
                # Los 41 intermedios ya cumplieron: sin ellos la carpeta queda
                # solo con lo que hay que mirar y lo que se entrega.
                shutil.rmtree(TEMPORAL, ignore_errors=True)
                self.cola.put(("listo", str(BASE / f"{base}.xlsm")))
            else:
                self.cola.put(("aviso",
                               "No se pudo armar el .xlsm con macros (hace falta "
                               f"Excel en esta maquina). Los 41 archivos sueltos, "
                               f"con sus botones, quedaron en:\n{TEMPORAL}"))
                self.cola.put(("listo", str(BASE / f"{base}.xlsx")))
        except Exception as e:                       # noqa: BLE001
            self.cola.put(("error", f"{type(e).__name__}: {e}"))
            self.cola.put(("fin", "Se corto por un error."))

    def _corre(self, orden):
        p = subprocess.Popen(orden, stdout=subprocess.PIPE,
                             stderr=subprocess.STDOUT, text=True,
                             encoding="utf-8", errors="replace", cwd=str(BASE),
                             creationflags=SIN_CONSOLA)
        for linea in p.stdout:
            self.cola.put(("linea", linea.rstrip()))
        return p.wait()

    # ------------------------------------------------------------ recepcion
    def _drena(self):
        """Pasa a la ventana lo que el hilo fue dejando en la cola."""
        try:
            while True:
                tipo, dato = self.cola.get_nowait()
                if tipo == "linea":
                    self._escribe(dato)
                    self._avance(dato)
                elif tipo == "paso":
                    self.estado.set(dato)
                elif tipo == "revisar":
                    self._carga_revisar(Path(dato))
                elif tipo == "aviso":
                    messagebox.showwarning("Sin macros", dato)
                elif tipo == "error":
                    self._escribe(dato, "err")
                elif tipo == "listo":
                    self.salida_final = Path(dato)
                    self.barra["value"] = 100
                    self.estado.set(f"Listo: {self.salida_final.name}")
                    self._escribe(f"\nListo: {dato}", "ok")
                    self.b_abrir.config(state="normal")
                elif tipo == "fin":
                    self.estado.set(dato)
                    self.corriendo = False
                    self.b_simular.config(state="normal")
                    self.b_generar.config(state="normal")
        except queue.Empty:
            pass
        self.after(80, self._drena)

    def _escribe(self, texto, tag=None):
        self.log.config(state="normal")
        self.log.insert("end", texto + "\n", tag or ())
        self.log.see("end")
        self.log.config(state="disabled")

    def _avance(self, linea):
        m = TOTAL.match(linea)
        if m:
            self.total = int(m.group(1))
            return
        if self.total and AVANCE.match(linea):
            self.hechos += 1
            self.barra["value"] = min(95, self.hechos * 90 / self.total)
            self.estado.set(f"Procesando… {self.hechos} de {self.total} plantillas")

    def _carga_revisar(self, ruta):
        self.tabla.delete(*self.tabla.get_children())
        if not ruta.exists():
            self.tabla.insert("", "end", values=(
                "Nada que revisar", "", "Todo calzo sin suponer nada.", "", ""))
            return
        try:
            from openpyxl import load_workbook
            hoja = load_workbook(ruta, read_only=True).active
            for i, fila in enumerate(hoja.iter_rows(values_only=True)):
                if i:
                    self.tabla.insert("", "end", values=fila)
        except Exception as e:                       # noqa: BLE001
            self._escribe(f"No se pudo leer {ruta.name}: {e}", "err")

    def abrir(self):
        if not self.salida_final or not self.salida_final.exists():
            return
        if os.name == "nt":
            os.startfile(self.salida_final)          # noqa: S606
        else:
            subprocess.Popen(["xdg-open", str(self.salida_final)])


if __name__ == "__main__":
    Interfaz().mainloop()
